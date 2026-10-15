"""Candidate management domain service classes and backward-compatible function interface.

This module houses all candidate domain service classes:
- BaseService: Safe database session management and common errors
- CandidateProfileService: Candidate listing, sorting, pagination, and details retrieval
- CandidateSearchService: Structured search, natural language search (LLM), filtering, and export
- CandidateParserService: PDF/Docx text extraction, resume classification, information extraction, and DB storage
- CandidateDocumentService: Resume tailoring, PDF/Docx generation, file preview, and bulk downloads
- CandidateService: Unified facade combining all domain services into a single interface
"""

import base64
from contextlib import contextmanager
import datetime
import json
import logging
import mimetypes
import os
import re
import shutil
import time as _time
from typing import Any, Dict, Generator, List, Optional, Tuple, Union
from uuid import UUID, uuid4
import zipfile
from email.utils import parseaddr

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt
from fastapi import HTTPException, status
from fastapi.responses import JSONResponse
from lxml import etree
import pandas as pd
from pydantic import ValidationError
from pypdf import PdfReader
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer
from sqlalchemy import JSON, Float, Integer, and_, case, cast, func, or_
from sqlalchemy.dialects.postgresql import aggregate_order_by
from sqlalchemy.orm import Session

from db.connection import SessionLocal
from src.admin.models import Admin, AiModelConfig
from src.candidate.models import (
    Attachment,
    Candidate,
    CandidateEducation,
    CandidateSkills,
    Company,
    Education,
    Resume,
    Role,
    Skill,
    WorkExperience,
)
from src.candidate.schema import (
    DynamicFilterRequest,
    DynamicFilterResponse,
    NLSearchPaginateRequest,
    NLSearchRequest,
    ResumeFilterRequest,
)
from src.client_track.models import CandidateInterviews
from src.email_reader.models import EmailLogs
from src.integrations.attachment_storage import (
    AttachmentStorageError,
    attachment_storage_service,
)
from src.integrations.llm import llm_service
from src.utils.response import serialize_response
from src.utils.helper import (
    _match_single_skill,
    _normalize_role_tokens,
    calculate_match_score,
    clean_mobile_number,
    compress_file,
    decrypt_data,
    record_model_usage,
    select_available_model,
)
from src.utils.redis_client import NL_SEARCH_TTL, get_redis

try:
    from openai import OpenAI
except ImportError:
    OpenAI = None

try:
    from anthropic import Anthropic
except ImportError:
    Anthropic = None

try:
    import ollama
except ImportError:
    ollama = None

logger = logging.getLogger(__name__)

# Constants
BASE_DIR = os.environ.get("BASE_DIR", "/app")
EXPORT_PATH = os.path.join(BASE_DIR, "export_files")
VALID_SORT_FIELDS = {"name", "experience", "year", "percentage"}
UPLOAD_DIR = os.path.join(BASE_DIR, "attachments")
FILES_DIR = "attachments/"
DOWNLOAD_FILES = "downloads/"
REDIS_KEY_PREFIX = "search:"

KEYWORD_STOP_WORDS = {
    "a", "an", "and", "as", "at", "candidate", "candidates", "developer",
    "developers", "engineer", "engineers", "for", "from", "in", "of", "or",
    "profile", "profiles", "resume", "resumes", "the", "with",
}


# =============================================================================
# Helper Utilities
# =============================================================================

def normalize_experience(exp_value: Any) -> int:
    """Ensure experience is always integer for DB and comparison."""
    if exp_value is None:
        return 0
    if isinstance(exp_value, (int, float)):
        return int(exp_value)
    text = str(exp_value).lower()
    match = re.search(r"\d+(\.\d+)?", text)
    if match:
        return int(float(match.group()))
    return 0


def parse_email_address(sender: str) -> str:
    """Extract lowercase email address from email header string."""
    _, addr = parseaddr(sender or "")
    return addr.strip().lower()


def clean_text(value: Any) -> str:
    """Sanitize string values for docx/pdf generation."""
    if value is None:
        return ""
    if isinstance(value, dict):
        return ""
    if isinstance(value, list):
        return ", ".join(map(str, value))
    return (
        str(value)
        .replace("\x00", "")
        .replace("fastapi_api |", "")
        .replace("postgres_db |", "")
        .strip()
    )


def add_heading(doc: Document, text: str):
    """Add a bold section heading to docx document."""
    heading = doc.add_heading(level=1)
    run = heading.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    return heading


# =============================================================================
# 1. Base Service
# =============================================================================

class CandidateServiceError(Exception):
    """Raised when candidate service encounters a recoverable business error."""

    def __init__(self, message: str, status_code: int = 400):
        self.message = message
        self.status_code = status_code
        super().__init__(self.message)


class BaseService:
    """Base class providing transactional DB session management."""

    def __init__(self, db: Optional[Session] = None):
        self._db = db

    @contextmanager
    def get_session(self, session: Optional[Session] = None) -> Generator[Session, None, None]:
        """Provide a transactional database session scope.
        
        If a session is explicitly passed or set on self._db, it is reused.
        Otherwise, a new SessionLocal() is created and cleanly closed on exit.
        """
        if session is not None:
            yield session
            return

        if self._db is not None:
            yield self._db
            return

        db = SessionLocal()
        try:
            yield db
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()


# =============================================================================
# 2. Candidate Profile Service
# =============================================================================

class CandidateProfileService(BaseService):
    """Handles candidate profile details, sorting, filtering and basic info retrieval."""

    def _validate_data(self, data: Any) -> None:
        """Raise if the query returned no rows."""
        if not data:
            raise CandidateServiceError("No data found", status_code=404)

    def _build_subqueries(self, db: Session):
        """Build reusable subqueries for candidate education, skills, and work experience."""
        candidate_education = (
            db.query(
                CandidateEducation.candidate_id.label("candidate_id"),
                func.json_agg(
                    func.json_build_object(
                        "education_id", Education.education_id,
                        "education", Education.education,
                        "institution", CandidateEducation.institution,
                        "percentage", CandidateEducation.percentage,
                        "year_of_passed", CandidateEducation.year_of_passed,
                    )
                ).label("education"),
            )
            .select_from(CandidateEducation)
            .join(Education, CandidateEducation.education_id == Education.education_id)
            .group_by(CandidateEducation.candidate_id)
            .subquery()
        )

        candidate_skills = (
            db.query(
                CandidateSkills.candidate_id.label("candidate_id"),
                func.json_agg(
                    func.json_build_object(
                        "skill_id", Skill.skill_id,
                        "skill", Skill.skill,
                    )
                ).label("skills"),
            )
            .select_from(CandidateSkills)
            .join(Skill, CandidateSkills.skill_id == Skill.skill_id)
            .group_by(CandidateSkills.candidate_id)
            .subquery()
        )

        candidate_work_exp = (
            db.query(
                WorkExperience.candidate_id.label("candidate_id"),
                func.json_agg(
                    func.json_build_object(
                        "role_id", Role.role_id,
                        "role", Role.role,
                        "company_name", Company.company_name,
                        "company_location", Company.company_location,
                        "start_date", WorkExperience.start_date,
                        "end_date", WorkExperience.end_date,
                        "is_present", WorkExperience.is_active,
                    )
                ).label("work_experience"),
            )
            .select_from(WorkExperience)
            .join(Role, WorkExperience.role_id == Role.role_id)
            .join(Company, WorkExperience.company_id == Company.company_id)
            .group_by(WorkExperience.candidate_id)
            .subquery()
        )

        return candidate_education, candidate_skills, candidate_work_exp

    def _base_candidate_query(self, db: Session, edu_sq, skills_sq, work_sq):
        """Base query selecting candidates with aggregated education, skills, and work experience."""
        return (
            db.query(
                func.json_build_object(
                    "candidate_id", Candidate.candidate_id,
                    "name", Candidate.name,
                    "email", Candidate.email_address,
                    "phone_number", Candidate.phone_number,
                    "location", Candidate.location,
                    "matching_role", Resume.candidate_role,
                    "total_experience", Candidate.total_experience,
                    "education", func.coalesce(edu_sq.c.education, cast("[]", JSON)),
                    "skills", func.coalesce(skills_sq.c.skills, cast("[]", JSON)),
                    "work_experience", func.coalesce(work_sq.c.work_experience, cast("[]", JSON)),
                ).label("candidate_info")
            )
            .select_from(Candidate)
            .outerjoin(edu_sq, Candidate.candidate_id == edu_sq.c.candidate_id)
            .outerjoin(skills_sq, Candidate.candidate_id == skills_sq.c.candidate_id)
            .outerjoin(work_sq, Candidate.candidate_id == work_sq.c.candidate_id)
            .outerjoin(Resume, Candidate.candidate_id == Resume.candidate_id)
            .filter(Candidate.is_active == True)
        )

    def _extract_results(self, rows, total_count: int) -> List[Dict[str, Any]]:
        """Unwrap SQLAlchemy rows and append total record count."""
        results = [row.candidate_info for row in rows]
        results.append({"total_record": total_count})
        return results

    def get_candidate_details(
        self, page: int = 1, sort_by: Optional[str] = None, sort_type: Optional[str] = None
    ) -> List[Dict[str, Any]]:
        """Return paginated candidate details with optional sorting."""
        with self.get_session() as db:
            try:
                per_page = 10
                offset = (int(page) - 1) * per_page

                candidate_education, candidate_skills, candidate_work_exp = self._build_subqueries(db)

                total_count = (
                    db.query(func.count(Candidate.candidate_id))
                    .filter(Candidate.is_active == True)
                    .scalar()
                ) or 0

                if total_count < offset and total_count > 0:
                    raise CandidateServiceError("Invalid page number", status_code=400)

                sort_type_resolved = sort_type if sort_type in ("asc", "desc") else "asc"

                if sort_by in (None, "None", ""):
                    query = (
                        self._base_candidate_query(db, candidate_education, candidate_skills, candidate_work_exp)
                        .limit(per_page)
                        .offset(offset)
                    )
                    data = query.all()
                    self._validate_data(data)
                    return self._extract_results(data, total_count)

                if sort_by == "name":
                    return self._sort_by_name(
                        db, sort_type_resolved, candidate_education, candidate_skills,
                        candidate_work_exp, per_page, offset, total_count
                    )

                if sort_by == "experience":
                    return self._sort_by_experience(
                        db, sort_type_resolved, candidate_education, candidate_skills,
                        candidate_work_exp, per_page, offset, total_count
                    )

                if sort_by == "year":
                    return self._sort_by_year(
                        db, sort_type_resolved, candidate_skills, candidate_work_exp,
                        per_page, offset, total_count
                    )

                if sort_by == "percentage":
                    return self._sort_by_percentage(
                        db, sort_type_resolved, candidate_education, candidate_skills,
                        candidate_work_exp, per_page, offset, total_count
                    )

                raise CandidateServiceError("Invalid sorting operation", status_code=400)

            except CandidateServiceError:
                raise
            except Exception as e:
                logger.error("[ERROR] in candidate fetch: %s", str(e), exc_info=True)
                raise CandidateServiceError(f"Failed to fetch candidates: {str(e)}", status_code=500)

    # Backward-compatible alias
    candidate_datails = get_candidate_details

    def _sort_by_name(self, db, sort_type, candidate_education, candidate_skills, candidate_work_exp, per_page, offset, total_count):
        order = Candidate.name.desc() if sort_type == "desc" else Candidate.name.asc()
        data = (
            self._base_candidate_query(db, candidate_education, candidate_skills, candidate_work_exp)
            .order_by(order)
            .limit(per_page)
            .offset(offset)
            .all()
        )
        self._validate_data(data)
        return self._extract_results(data, total_count)

    def _sort_by_experience(self, db, sort_type, candidate_education, candidate_skills, candidate_work_exp, per_page, offset, total_count):
        order = Candidate.total_experience.desc() if sort_type == "desc" else Candidate.total_experience.asc()
        data = (
            self._base_candidate_query(db, candidate_education, candidate_skills, candidate_work_exp)
            .order_by(order)
            .limit(per_page)
            .offset(offset)
            .all()
        )
        self._validate_data(data)
        return self._extract_results(data, total_count)

    def _sort_by_year(self, db, sort_type, candidate_skills, candidate_work_exp, per_page, offset, total_count):
        agg_func = func.max if sort_type == "desc" else func.min
        agg_order = CandidateEducation.year_of_passed.desc() if sort_type == "desc" else CandidateEducation.year_of_passed.asc()
        sort_order_col_dir = "desc" if sort_type == "desc" else "asc"

        passout_year_sq = (
            db.query(
                CandidateEducation.candidate_id.label("candidate_id"),
                agg_func(CandidateEducation.year_of_passed).label("latest_year"),
                func.json_agg(
                    aggregate_order_by(
                        func.json_build_object(
                            "education_id", Education.education_id,
                            "education", Education.education,
                            "institution", CandidateEducation.institution,
                            "percentage", CandidateEducation.percentage,
                            "year_of_passed", CandidateEducation.year_of_passed,
                        ),
                        agg_order,
                    )
                ).label("education"),
            )
            .join(Education, CandidateEducation.education_id == Education.education_id)
            .filter(~func.lower(Education.education).in_(["sslc", "hse"]))
            .group_by(CandidateEducation.candidate_id)
            .subquery()
        )

        order_col = passout_year_sq.c.latest_year.desc() if sort_order_col_dir == "desc" else passout_year_sq.c.latest_year.asc()

        data = (
            db.query(
                func.json_build_object(
                    "candidate_id", Candidate.candidate_id,
                    "name", Candidate.name,
                    "email", Candidate.email_address,
                    "phone_number", Candidate.phone_number,
                    "location", Candidate.location,
                    "total_experience", Candidate.total_experience,
                    "education", func.coalesce(passout_year_sq.c.education, cast("[]", JSON)),
                    "skills", func.coalesce(candidate_skills.c.skills, cast("[]", JSON)),
                    "work_experience", func.coalesce(candidate_work_exp.c.work_experience, cast("[]", JSON)),
                ).label("candidate_info")
            )
            .select_from(Candidate)
            .outerjoin(passout_year_sq, Candidate.candidate_id == passout_year_sq.c.candidate_id)
            .outerjoin(candidate_skills, Candidate.candidate_id == candidate_skills.c.candidate_id)
            .outerjoin(candidate_work_exp, Candidate.candidate_id == candidate_work_exp.c.candidate_id)
            .filter(Candidate.is_active == True)
            .order_by(order_col)
            .limit(per_page)
            .offset(offset)
            .all()
        )
        self._validate_data(data)
        return self._extract_results(data, total_count)

    def _sort_by_percentage(self, db, sort_type, candidate_education, candidate_skills, candidate_work_exp, per_page, offset, total_count):
        agg_func = func.max if sort_type == "desc" else func.min

        percentage_sq = (
            db.query(
                CandidateEducation.candidate_id.label("candidate_id"),
                agg_func(CandidateEducation.percentage).label("highest_percentage"),
            )
            .group_by(CandidateEducation.candidate_id)
            .subquery()
        )

        order = percentage_sq.c.highest_percentage.desc() if sort_type == "desc" else percentage_sq.c.highest_percentage.asc()

        data = (
            self._base_candidate_query(db, candidate_education, candidate_skills, candidate_work_exp)
            .outerjoin(percentage_sq, Candidate.candidate_id == percentage_sq.c.candidate_id)
            .order_by(order)
            .limit(per_page)
            .offset(offset)
            .all()
        )
        self._validate_data(data)
        return self._extract_results(data, total_count)

    def get_candidate_basic_info(
        self,
        filter_by: Optional[str] = None,
        filter_value: Optional[str] = None,
        page: int = 1,
        page_size: int = 10,
        sort_by: Optional[str] = None,
        sort_order: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Fetch simplified active candidate information."""
        with self.get_session() as db:
            try:
                sort_map = {
                    "candidate_name": Candidate.name,
                    "email": Candidate.email_address,
                    "phone_number": Candidate.phone_number,
                    "total_experience": Candidate.total_experience,
                    "role": Resume.candidate_role,
                }

                filter_map = {
                    "candidate_name": (Candidate.name, "string"),
                    "name": (Candidate.name, "string"),
                    "email": (Candidate.email_address, "string"),
                    "phone_number": (Candidate.phone_number, "string"),
                    "total_experience": (Candidate.total_experience, "integer"),
                    "role": (Resume.candidate_role, "string"),
                }

                query = (
                    db.query(
                        Candidate.candidate_id,
                        Candidate.name.label("candidate_name"),
                        Resume.resume_id,
                        Resume.candidate_role.label("candidate_role"),
                        Candidate.email_address,
                        Candidate.phone_number,
                        Candidate.total_experience,
                    )
                    .join(Resume, Candidate.candidate_id == Resume.candidate_id)
                    .filter(Candidate.is_active == True)
                )

                sort_order_dir = sort_order if sort_order in ("asc", "desc") else "asc"
                if sort_by in sort_map:
                    col = sort_map[sort_by]
                    query = query.order_by(col.desc() if sort_order_dir == "desc" else col.asc())

                if filter_by in filter_map and filter_value:
                    col, col_type = filter_map[filter_by]
                    if col_type == "string":
                        query = query.filter(col.ilike(f"%{filter_value}%"))
                    elif col_type == "integer":
                        query = query.filter(col == int(filter_value))

                page_int = int(page)
                page_size_int = int(page_size)
                if page_int > 0 and page_size_int > 0:
                    query = query.limit(page_size_int).offset((page_int - 1) * page_size_int)

                data = query.all()
                data_serialized = [
                    {
                        "candidate_id": row.candidate_id,
                        "candidate_name": row.candidate_name,
                        "resume_id": row.resume_id,
                        "role": row.candidate_role,
                        "email": row.email_address,
                        "phone_number": row.phone_number,
                        "total_experience": row.total_experience,
                    }
                    for row in data
                ]

                return {
                    "status": status.HTTP_200_OK,
                    "message": "Candidates basic information retrieved successfully",
                    "data": data_serialized,
                }
            except HTTPException:
                raise
            except Exception as e:
                logger.warning("[get_candidate_basic_info] Error: %s", str(e), exc_info=True)
                raise HTTPException(
                    status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                    detail={
                        "status": "error",
                        "message": str(e),
                    },
                )


# =============================================================================
# 3. Candidate Search Service
# =============================================================================

class CandidateSearchService(BaseService):
    """Encapsulates structured resume search, natural language queries, and exports."""

    @staticmethod
    def _build_redis_key(search_id: str) -> str:
        return f"{REDIS_KEY_PREFIX}{search_id}"

    def store_search_session(
        self,
        search_id: str,
        raw_filters: dict,
        resolved_filters: dict,
        user_query: str,
    ) -> None:
        """Persist extracted + resolved filters in Redis with TTL."""
        r = get_redis()
        payload = {
            "filters": resolved_filters,
            "raw_filters": raw_filters,
            "user_query": user_query,
            "created_at": datetime.datetime.utcnow().isoformat(),
        }
        r.setex(self._build_redis_key(search_id), NL_SEARCH_TTL, json.dumps(payload))
        logger.info("[nl_search] Stored session %s (TTL=%ds)", search_id, NL_SEARCH_TTL)

    def load_search_session(self, search_id: str) -> Optional[dict]:
        """Retrieve a cached search session, or None if missing."""
        r = get_redis()
        raw = r.get(self._build_redis_key(search_id))
        if raw is None:
            return None
        try:
            return json.loads(raw)
        except (json.JSONDecodeError, TypeError):
            logger.warning("[nl_search] Corrupt cache entry for %s", search_id)
            return None

    def _keyword_filter_attempts(self, user_query: str) -> List[dict]:
        """Build deterministic fallback filters when LLM extraction is unavailable."""
        cleaned = re.sub(r"\s+", " ", user_query or "").strip()
        if not cleaned:
            return []

        tokens = [
            token
            for token in re.findall(r"[A-Za-z0-9+#.]+", cleaned)
            if len(token) > 1 and token.lower() not in KEYWORD_STOP_WORDS
        ]
        terms = list(dict.fromkeys([cleaned] + tokens))

        attempts: List[dict] = []
        for field in ("skills", "roles", "companies", "education"):
            attempts.extend({field: [term]} for term in terms)
        attempts.append({"name": cleaned})
        return attempts

    def _fallback_extract_and_resolve(self, user_query: str) -> Tuple[dict, dict]:
        """Try keyword-based filters so simple searches work without LLM.

        Strategy:
        1. Keyword-match skills/education from tokens in the query.
        2. Apply _complete_query_filters to extract experience ranges and
           job-title roles from the raw query text via regex.
        3. Merge both sources, then resolve to UUIDs.
        """
        # Step 1: collect only skill and education matches via keyword search
        # (avoid adding roles/companies from token matching — they produce false positives)
        keyword_raw: dict = {}
        for raw_filters in self._keyword_filter_attempts(user_query):
            for field in ("skills", "education"):
                if field in raw_filters:
                    resolved_partial = self.resolve_dynamic_filters({field: raw_filters[field]})
                    if resolved_partial.get(field):
                        existing = keyword_raw.get(field, [])
                        for item in raw_filters[field]:
                            if item not in existing:
                                existing.append(item)
                        keyword_raw[field] = existing

        # Step 2: enrich with experience + role via regex on the query text
        enriched_raw = self._complete_query_filters(keyword_raw, user_query)

        # If no skills/roles/education/experience found, check if query represents a candidate name
        if (
            not enriched_raw.get("skills")
            and not enriched_raw.get("roles")
            and not enriched_raw.get("education")
            and not enriched_raw.get("min_experience")
            and not enriched_raw.get("name")
        ):
            cleaned_name = re.sub(
                r"^(?:find|search|show|get|view|filter)?\s*(?:candidate|resumes?|profiles?)?\s*(?:named?|name:?)?\s*",
                "",
                user_query or "",
                flags=re.IGNORECASE,
            ).strip()
            if len(cleaned_name) >= 2 and re.search(r"[A-Za-z]", cleaned_name):
                enriched_raw["name"] = [cleaned_name]

        if enriched_raw:
            resolved = self.resolve_dynamic_filters(enriched_raw)
            if resolved:
                logger.info(
                    "[nl_search] Keyword fallback matched: %s -> %s",
                    enriched_raw, resolved,
                )
                return enriched_raw, resolved

        # Step 3: last-resort — return scalar-only fields even without UUID resolution
        scalar_only: dict = {}
        for k in ("min_experience", "max_experience", "name"):
            if enriched_raw.get(k) is not None:
                scalar_only[k] = enriched_raw[k]
        if enriched_raw.get("roles"):
            scalar_only["roles"] = enriched_raw["roles"]
        if enriched_raw.get("skills"):
            scalar_only["skills"] = enriched_raw["skills"]

        if scalar_only:
            resolved = self.resolve_dynamic_filters(scalar_only) or scalar_only
            logger.info(
                "[nl_search] Keyword fallback (scalar only): %s -> %s",
                scalar_only, resolved,
            )
            return scalar_only, resolved

        return {}, {}

    @staticmethod
    def _complete_query_filters(extracted: dict, query: str) -> dict:
        """Normalize LLM keys and recover basic criteria omitted by an LLM."""
        aliases = {
            "skills": "skills", "education": "education", "roles": "roles",
            "companies": "companies", "minexperience": "min_experience",
            "maxexperience": "max_experience", "name": "name",
            "candidatename": "name",
            "passoutstartyear": "passout_start_year",
            "passoutendyear": "passout_end_year", "percentage": "percentage",
        }
        result: dict = {}
        for key, value in (extracted or {}).items():
            normalized_key = re.sub(r"[^a-z]", "", str(key).lower())
            canonical_key = aliases.get(normalized_key)
            if canonical_key and value is not None:
                result[canonical_key] = value
            elif normalized_key in ("experience", "totalexperience", "exp", "yearsofexperience") and value is not None:
                if isinstance(value, (int, float)):
                    if result.get("min_experience") is None:
                        result["min_experience"] = float(value)
                elif isinstance(value, str):
                    range_m = re.search(r"(\d+(?:\.\d+)?)\s*(?:-|to)\s*(\d+(?:\.\d+)?)", value)
                    if range_m:
                        if result.get("min_experience") is None:
                            result["min_experience"] = float(range_m.group(1))
                        if result.get("max_experience") is None:
                            result["max_experience"] = float(range_m.group(2))
                    else:
                        num_m = re.search(r"(\d+(?:\.\d+)?)", value)
                        if num_m and result.get("min_experience") is None:
                            result["min_experience"] = float(num_m.group(1))

        # Ensure min_experience / max_experience are floats if present
        for exp_key in ("min_experience", "max_experience"):
            val = result.get(exp_key)
            if val is not None:
                if isinstance(val, (int, float)):
                    result[exp_key] = float(val)
                elif isinstance(val, str):
                    m = re.search(r"(\d+(?:\.\d+)?)", val)
                    result[exp_key] = float(m.group(1)) if m else None
                else:
                    result[exp_key] = None

        for field in ("skills", "education", "roles", "companies"):
            if result.get(field) is not None and not isinstance(result[field], list):
                result[field] = [result[field]]

        if result.get("name") is not None:
            if isinstance(result["name"], list):
                cleaned_names = [str(n).strip() for n in result["name"] if str(n).strip()]
                result["name"] = cleaned_names if cleaned_names else None
            elif isinstance(result["name"], str):
                cleaned_name = result["name"].strip()
                result["name"] = [cleaned_name] if cleaned_name else None
            else:
                result["name"] = None

        # Remove experience-like terms from skills if accidentally extracted
        exp_patterns = [
            r'\b\d+\s*(?:-|to)\s*\d+\s*(?:years?|yrs?|exp|yoe)',
            r'\b\d+\+\s*(?:years?|yrs?|exp|yoe)(?:\s+of\s+experience|\s+exp)?',
            r'\b\d+\s*(?:years?|yrs?|exp|yoe)(?:\s+of\s+experience|\s+exp)?\b(?!\s*\+)',
        ]
        if result.get("skills"):
            cleaned_skills = []
            for skill in result["skills"]:
                s = str(skill).lower()
                is_exp = False
                for pat in exp_patterns:
                    if re.search(pat, s, re.IGNORECASE):
                        is_exp = True
                        break
                if not is_exp:
                    cleaned_skills.append(skill)
            if cleaned_skills != result["skills"]:
                result["skills"] = cleaned_skills

        text = query or ""
        range_match = re.search(
            r"\b(?:between\s+)?(\d+(?:\.\d+)?)\s*(?:-|to)\s*(\d+(?:\.\d+)?)\s*(?:years?|yrs?|yoe|exp|experience)?\b",
            text,
            re.IGNORECASE,
        )
        if range_match:
            if result.get("min_experience") is None:
                result["min_experience"] = float(range_match.group(1))
            if result.get("max_experience") is None:
                result["max_experience"] = float(range_match.group(2))
        else:
            max_match = re.search(
                r"\b(?:max|maximum|up\s+to|less\s+than|under|below|\<=?)\s*(\d+(?:\.\d+)?)\s*(?:years?|yrs?|yoe|exp|experience)?\b",
                text,
                re.IGNORECASE,
            )
            if max_match:
                if result.get("max_experience") is None:
                    result["max_experience"] = float(max_match.group(1))
            else:
                if result.get("min_experience") is None:
                    min_plus_match = re.search(
                        r"\b(\d+(?:\.\d+)?)\s*\+\s*(?:years?|yrs?|yoe|exp|experience)?\b",
                        text,
                        re.IGNORECASE,
                    )
                    min_word_match = re.search(
                        r"\b(?:min|minimum|at\s+least|more\s+than|greater\s+than|above|\>=?)\s*(\d+(?:\.\d+)?)\s*(?:years?|yrs?|yoe|exp|experience)?\b",
                        text,
                        re.IGNORECASE,
                    )
                    single_match = re.search(
                        r"\b(\d+(?:\.\d+)?)\s*(?:years?|yrs?|yoe)(?:\s+(?:of\s+)?(?:experience|exp))?\b",
                        text,
                        re.IGNORECASE,
                    )
                    exp_prefix_match = re.search(
                        r"\b(?:experience|exp|total\s+experience)\s*(?:of|:|\s)\s*(\d+(?:\.\d+)?)\b",
                        text,
                        re.IGNORECASE,
                    )
                    if min_plus_match:
                        result["min_experience"] = float(min_plus_match.group(1))
                    elif min_word_match:
                        result["min_experience"] = float(min_word_match.group(1))
                    elif single_match:
                        result["min_experience"] = float(single_match.group(1))
                    elif exp_prefix_match:
                        result["min_experience"] = float(exp_prefix_match.group(1))

        # Detect role title from query text
        role_noun_pattern = r"\b(?:developer|engineer|architect|analyst|designer|manager|consultant|administrator|specialist|programmer|scientist|tester|lead)\b"
        role_title_pattern = r"\b([A-Za-z][A-Za-z0-9+#. /-]{0,50}?" + role_noun_pattern + r")\b"
        detected_role = None
        role_match = re.search(role_title_pattern, text, re.IGNORECASE)
        if role_match:
            detected_role = re.sub(r"\s+", " ", role_match.group(1)).strip().title()

        roles = [str(r).strip() for r in (result.get("roles") or []) if str(r).strip()]

        # If LLM classified a job title as a skill, move it to roles
        raw_skills = [str(s).strip() for s in (result.get("skills") or []) if str(s).strip()]
        non_role_skills = []
        for s in raw_skills:
            if re.search(role_noun_pattern + r"$", s, re.IGNORECASE):
                s_title = s.title()
                if s_title not in roles:
                    roles.append(s_title)
            else:
                non_role_skills.append(s)

        if not roles and detected_role:
            roles = [detected_role]
        elif detected_role:
            # Upgrade generic single word role (e.g. 'Analyst') to full title ('Data Analyst') if detected in query
            upgraded = []
            for r in roles:
                if r.lower() in detected_role.lower() and len(r) < len(detected_role):
                    upgraded.append(detected_role)
                else:
                    upgraded.append(r)
            roles = list(dict.fromkeys(upgraded))

        if not result.get("name"):
            name_match = re.search(
                r"\b(?:candidate\s+name|candidate\s+named|candidate|name\s+is|named|name|resume\s+of|profile\s+of)\s*[:=]?\s*([A-Za-z][A-Za-z.'-]{1,25}(?:\s+[A-Za-z][A-Za-z.'-]{1,25})?)\b",
                text,
                re.IGNORECASE,
            )
            if name_match:
                potential_name = name_match.group(1).strip()
                if (
                    potential_name.lower() not in KEYWORD_STOP_WORDS
                    and not re.search(role_noun_pattern, potential_name, re.IGNORECASE)
                ):
                    result["name"] = [potential_name]

        # Clean experience, candidate, and name prefixes from roles
        extracted_names = result.get("name") or []
        cleaned_roles = []
        for r in roles:
            clean_r = re.sub(
                r"^(?:\d+\+?\s*|\b(?:years?|yrs?|yoe|experience|exp|of|candidates?|resumes?|profiles?)\b\s*)+",
                "",
                r,
                flags=re.IGNORECASE,
            ).strip()
            for cand_n in extracted_names:
                clean_r = re.sub(r"\b" + re.escape(str(cand_n)) + r"\b", "", clean_r, flags=re.IGNORECASE).strip()
            if clean_r:
                cleaned_roles.append(clean_r)
        roles = list(dict.fromkeys(cleaned_roles))

        # Filter out role tokens from skills unless explicitly requested as skills
        skill_clause_match = re.search(
            r"\b(?:with|skills?\s*(?:in|of|:)?|having|knowing)\s+([a-zA-Z0-9+#,\s/]+)",
            text, re.IGNORECASE,
        )
        skill_clause_text = skill_clause_match.group(1).lower() if skill_clause_match else ""
        skill_clause_text = re.sub(r"\b\d+\+?\s*(?:years?|yrs?)[^,]*", "", skill_clause_text)

        role_tokens = set()
        for r in roles:
            for w in re.findall(r"[a-zA-Z0-9+#]+", r.lower()):
                role_tokens.add(w)

        final_skills = []
        for s in non_role_skills:
            s_lower = s.lower().strip()
            s_tokens = set(re.findall(r"[a-zA-Z0-9+#]+", s_lower))

            # If skill consists of role tokens (e.g. 'data' or 'analyst' from 'Data Analyst')
            if s_tokens.issubset(role_tokens):
                if not skill_clause_text or not re.search(r"\b" + re.escape(s_lower) + r"\b", skill_clause_text):
                    continue
            final_skills.append(s)

        result["roles"] = roles
        result["skills"] = final_skills
        return result

    def extract_filters_from_query(self, query: str) -> dict:
        """Use LLM to convert a natural-language query into structured filters."""
        prompt = """Extract recruiter-search filters. Return JSON only with exactly these keys:
{"skills":[],"education":[],"roles":[],"companies":[],"min_experience":null,"max_experience":null,"name":null,"passout_start_year":null,"passout_end_year":null,"percentage":null}

CRITICAL INSTRUCTIONS:
1. ROLES VS SKILLS:
   - If the user query is a job role/title (e.g. "data analyst", "python developer", "frontend engineer"), extract it under "roles" ONLY (e.g. roles: ["Data Analyst"]).
   - DO NOT put words from the role title (such as "data", "analyst", "python", "developer") into "skills".
   - ONLY extract "skills" if the query explicitly mentions skills, tools, or technologies separate from the role title (e.g. "data analyst with python and sql skills", "skills in react", "python developer with django").
2. EXPERIENCE:
   - If the query mentions experience (e.g. "3 years", "3+ years", "3-5 years", "3 to 5 years", "minimum 3 years", "3 yrs exp", "4 years of experience", "up to 5 years"):
     - For single value, minimum, or "+": set "min_experience" to the number as integer or float (e.g. "3+ years" -> min_experience: 3).
     - For ranges: set both "min_experience" and "max_experience" (e.g. "3-5 years" -> min_experience: 3, max_experience: 5).
     - For maximum: set "max_experience" (e.g. "up to 5 years" -> max_experience: 5).
     - Always return experience as numeric values (e.g. 3, not "3" or "3+").
3. CANDIDATE NAME:
   - If the query searches for a person/candidate by name (e.g. "John Doe", "candidate Aarav", "find Alice Smith", "John"), extract it under "name" (e.g. name: "John Doe").
4. Examples:
   - "data analyst" -> {"roles": ["Data Analyst"], "skills": []}
   - "python developer 3+ years" -> {"roles": ["Python Developer"], "skills": [], "min_experience": 3}
   - "python developer with 3-5 years exp" -> {"roles": ["Python Developer"], "skills": [], "min_experience": 3, "max_experience": 5}
   - "data analyst with 3 years experience" -> {"roles": ["Data Analyst"], "skills": [], "min_experience": 3}
   - "data analyst with python and sql skills" -> {"roles": ["Data Analyst"], "skills": ["Python", "SQL"]}
   - "frontend developer with react, typescript, 4 years" -> {"roles": ["Frontend Developer"], "skills": ["React", "TypeScript"], "min_experience": 4}
   - "candidates with aws and docker skills" -> {"roles": [], "skills": ["AWS", "Docker"]}
   - "John Doe" -> {"name": "John Doe", "roles": [], "skills": []}
   - "candidate Aarav Sharma" -> {"name": "Aarav Sharma", "roles": [], "skills": []}
   - "John Doe python developer" -> {"name": "John Doe", "roles": ["Python Developer"], "skills": []}

Use [] or null when absent. Do not omit or add keys."""
        with self.get_session() as db:
            try:
                # Find an admin that has at least one active model config
                admin = (
                    db.query(Admin)
                    .join(AiModelConfig, AiModelConfig.admin_id == Admin.admin_id)
                    .filter(Admin.is_active == True, AiModelConfig.is_active == True)
                    .first()
                )
                if not admin:
                    # Fallback: any active admin (will trigger the "no model" error below)
                    admin = db.query(Admin).filter(Admin.is_active == True).first()
                if not admin:
                    logger.warning("[extract_filters_from_query] No admin found")
                    raise CandidateServiceError("No admin found", status_code=500)

                model_info = select_available_model(admin.admin_id)
                if not model_info:
                    raise CandidateServiceError("No active AI model configured", status_code=500)

                model_config_id = model_info.get("model_config_id")
                model = model_info.get("model_name", "ollama").lower()
                version = model_info.get("version_name", "llama3").lower()
                base_url = model_info.get("base_url")
                api_key = decrypt_data(model_info.get("apikey"))
                logger.info(
                    "[nl_search] Selected LLM configuration | provider=%s | model=%s | base_url=%s | config_id=%s",
                    model, version, base_url or "<provider-default>", model_config_id,
                )

                message = [
                    {
                        "role": "system",
                        "content": "You are a strict JSON extractor for recruiter search queries. Respond with a single valid JSON object only. No markdown fences, no commentary.",
                    },
                    {"role": "user", "content": prompt + "\n\nQuery:\n" + query[:1000]},
                ]

                if model == "openai":
                    if not api_key:
                        raise CandidateServiceError("OpenAI API key not found", status_code=500)
                    client = OpenAI(
                        api_key=api_key,
                        base_url=base_url or "https://api.openai.com/v1",
                    )
                    response = client.chat.completions.create(model=version, messages=message)
                    content = response.choices[0].message.content.strip()
                    token_used = response.usage.total_tokens
                    record_model_usage(model_config_id, token_used, admin.name)

                elif model == "groq":
                    if not api_key:
                        raise CandidateServiceError("Groq API key not found", status_code=500)
                    client = OpenAI(
                        api_key=api_key,
                        base_url=base_url or "https://api.groq.com/openai/v1",
                    )
                    response = client.chat.completions.create(
                        model=version,
                        messages=message,
                        response_format={"type": "json_object"},
                    )
                    content = response.choices[0].message.content.strip()
                    if response.usage:
                        record_model_usage(model_config_id, response.usage.total_tokens, admin.name)

                elif model == "claude":
                    if not api_key:
                        raise CandidateServiceError("Claude API key not found", status_code=500)
                    client = Anthropic(
                        api_key=api_key,
                        base_url=base_url or "https://api.anthropic.com",
                    )
                    response = client.messages.create(
                        model=version,
                        max_tokens=1000,
                        system="You are a strict JSON extractor for recruiter search queries. Respond with a single valid JSON object only. No markdown fences, no commentary.",
                        messages=[{"role": "user", "content": prompt + "\n\nQuery:\n" + query[:1000]}],
                    )
                    content = response.content[0].text.strip()

                else:
                    response = ollama.chat(model=version, messages=message)
                    content = response["message"]["content"].strip()

                start = content.find("{")
                end = content.rfind("}") + 1
                if start == -1 or end == 0:
                    return {}

                parsed = json.loads(content[start:end])
                if not isinstance(parsed, dict):
                    return {}

                logger.info(
                    "[nl_search] LLM extracted filters | query=%r | filters=%s",
                    query, parsed,
                )

                allowed_keys = {
                    "skills", "education", "roles", "companies",
                    "min_experience", "max_experience", "name",
                    "passout_start_year", "passout_end_year", "percentage",
                }
                normalized = self._complete_query_filters(parsed, query)
                extracted_filters = {
                    k: v for k, v in normalized.items()
                    if k in allowed_keys and v is not None
                }
                logger.info(
                    "[nl_search] Normalized filters | query=%r | filters=%s",
                    query, extracted_filters,
                )
                return extracted_filters

            except Exception as e:
                logger.error("[extract_filters_from_query] Extraction failed: %s", str(e), exc_info=True)
                return {}

    def resolve_dynamic_filters(self, dynamic_filters: dict) -> dict:
        """Convert human-readable names to UUID-based values matching database records.

        Skills and roles are resolved by fetching all DB rows once and matching
        in Python (alias-aware / fuzzy for skills; token-overlap for roles).
        The matched IDs are then used with an IN operator when filtering candidates.
        """
        if not dynamic_filters or not isinstance(dynamic_filters, dict):
            return {}

        with self.get_session() as db:
            resolved: dict = {}
            try:
                def _match_skill_ids(values: list) -> list:
                    """Fetch all skills once, match each query keyword using
                    alias-aware / fuzzy comparison, return matched UUIDs."""
                    if not values:
                        return []
                    all_skills = db.query(Skill.skill_id, Skill.skill).all()
                    matched_ids = []
                    for value in values:
                        if not value or not isinstance(value, str) or not value.strip():
                            continue
                        query_kw = value.strip()
                        for row in all_skills:
                            if _match_single_skill(query_kw, [row.skill]):
                                sid = str(row.skill_id)
                                if sid not in matched_ids:
                                    matched_ids.append(sid)
                    return matched_ids

                def _match_role_ids(values: list) -> list:
                    """Fetch all roles once, match each query keyword using
                    case-insensitive exact check then token-overlap, return matched UUIDs."""
                    if not values:
                        return []
                    all_roles = db.query(Role.role_id, Role.role).all()
                    matched_ids = []
                    for value in values:
                        if not value or not isinstance(value, str) or not value.strip():
                            continue
                        query_kw = value.strip().lower()
                        q_tokens = set(_normalize_role_tokens(query_kw))
                        for row in all_roles:
                            db_name = (row.role or "").strip().lower()
                            # exact case-insensitive match
                            if query_kw == db_name:
                                rid = str(row.role_id)
                                if rid not in matched_ids:
                                    matched_ids.append(rid)
                                continue
                            # token overlap: threshold scales with query length
                            #  1 token  → exact only (no overlap shortcut)
                            #  2 tokens → both must match (strict)
                            #  3+ tokens → at least ceil(2/3) must match
                            if len(q_tokens) < 2:
                                continue
                            db_tokens = set(_normalize_role_tokens(db_name))
                            if not db_tokens:
                                continue
                            overlap = q_tokens & db_tokens
                            if len(q_tokens) == 2:
                                required = 2           # both tokens must match
                            else:
                                required = max(2, -(-len(q_tokens) * 2 // 3))  # ceil(2/3)
                            if len(overlap) >= required:
                                rid = str(row.role_id)
                                if rid not in matched_ids:
                                    matched_ids.append(rid)
                    return matched_ids

                def _match_ids(model, id_attr, name_attr, values: list) -> list:
                    """Generic fallback: fetch all rows once, case-insensitive exact match
                    then IN operator — used for education and other simple tables."""
                    if not values:
                        return []
                    query_lower = [
                        v.strip().lower() for v in values
                        if v and isinstance(v, str) and v.strip()
                    ]
                    if not query_lower:
                        return []
                    all_rows = db.query(model).all()
                    ids = []
                    for row in all_rows:
                        db_val = (getattr(row, name_attr.key, None) or "").strip().lower()
                        if db_val in query_lower:
                            rid = str(getattr(row, id_attr))
                            if rid not in ids:
                                ids.append(rid)
                    return ids

                skill_names = dynamic_filters.get("skills")
                if skill_names and isinstance(skill_names, list):
                    skill_ids = _match_skill_ids(skill_names)
                    if skill_ids:
                        resolved["skills"] = skill_ids

                edu_names = dynamic_filters.get("education")
                if edu_names and isinstance(edu_names, list):
                    edu_ids = _match_ids(Education, "education_id", Education.education, edu_names)
                    if edu_ids:
                        resolved["education"] = edu_ids

                role_names = dynamic_filters.get("roles")
                if role_names and isinstance(role_names, list):
                    role_ids = _match_role_ids(role_names)
                    if role_ids:
                        resolved["roles"] = list(dict.fromkeys(role_ids))
                    else:
                        # Keep the raw title so search_resumes can also search
                        # Resume.candidate_role, which may not have a master Role row.
                        resolved["roles"] = [
                            role.strip() for role in role_names
                            if isinstance(role, str) and role.strip()
                        ]

                companies = dynamic_filters.get("companies")
                if companies and isinstance(companies, list):
                    resolved["companies"] = [c.strip() for c in companies if isinstance(c, str) and c.strip()]

                for field in ("min_experience", "max_experience"):
                    val = dynamic_filters.get(field)
                    if val is not None:
                        if isinstance(val, str):
                            m = re.search(r"(\d+(?:\.\d+)?)", val)
                            if m:
                                resolved[field] = float(m.group(1))
                        else:
                            try:
                                resolved[field] = float(val)
                            except (ValueError, TypeError):
                                pass

                name_val = dynamic_filters.get("name") or dynamic_filters.get("candidate_name")
                if name_val:
                    if isinstance(name_val, str) and name_val.strip():
                        resolved["name"] = [name_val.strip()]
                    elif isinstance(name_val, list):
                        cleaned_names = [str(n).strip() for n in name_val if str(n).strip()]
                        if cleaned_names:
                            resolved["name"] = cleaned_names

                for field in ("passout_start_year", "passout_end_year"):
                    val = dynamic_filters.get(field)
                    if val is not None:
                        if isinstance(val, str):
                            m = re.search(r"(\d{4})", val)
                            if m:
                                resolved[field] = int(m.group(1))
                        else:
                            try:
                                resolved[field] = int(val)
                            except (ValueError, TypeError):
                                pass

                pct_val = dynamic_filters.get("percentage")
                if pct_val is not None:
                    if isinstance(pct_val, str):
                        m = re.search(r"(\d+(?:\.\d+)?)", pct_val)
                        if m:
                            resolved["percentage"] = float(m.group(1))
                    else:
                        try:
                            resolved["percentage"] = float(pct_val)
                        except (ValueError, TypeError):
                            pass

                return resolved

            except Exception as e:
                logger.error("[resolve_dynamic_filters] Error: %s", str(e), exc_info=True)
                return {}

    def extract_and_resolve(self, user_query: str) -> Tuple[dict, dict]:
        """Run LLM extraction -> UUID resolution pipeline."""
        raw_filters = self.extract_filters_from_query(user_query)
        if not raw_filters:
            raw_filters, resolved = self._fallback_extract_and_resolve(user_query)
            if not resolved:
                raise ValueError("Could not extract any filters from the query. Please rephrase and try again.")
            return raw_filters, resolved

        try:
            DynamicFilterResponse(**raw_filters)
        except ValidationError:
            pass

        resolved = self.resolve_dynamic_filters(raw_filters)
        logger.info(
            "[nl_search] Resolved database filters | extracted=%s | resolved=%s",
            raw_filters, resolved,
        )
        if not resolved:
            fallback_raw, fallback_resolved = self._fallback_extract_and_resolve(user_query)
            if fallback_resolved:
                return fallback_raw, fallback_resolved
            raise ValueError("Candidate not found for the given query.")

        return raw_filters, resolved

    @staticmethod
    def merge_filters(standard: dict, dynamic: dict) -> dict:
        """Merge standard (UI-driven) filters with resolved dynamic (LLM-driven) filters."""
        if not dynamic:
            return dict(standard) if standard else {}
        if not standard:
            return dict(dynamic)

        merged = dict(standard)
        list_fields = ("skills", "education", "roles", "companies")
        for key in list_fields:
            std_vals = standard.get(key) or []
            dyn_vals = dynamic.get(key) or []
            if not isinstance(std_vals, list):
                std_vals = []
            if not isinstance(dyn_vals, list):
                dyn_vals = []
            combined = list(dict.fromkeys(std_vals + dyn_vals))
            if combined:
                merged[key] = combined

        scalar_fields = (
            "min_experience", "max_experience", "name", "percentage",
            "passout_start_year", "passout_end_year", "file_name",
        )
        for key in scalar_fields:
            if merged.get(key) is None and dynamic.get(key) is not None:
                merged[key] = dynamic[key]

        for key in list_fields:
            if key in merged and not merged[key]:
                del merged[key]

        return merged

    def search_resumes(self, filters: dict, export: bool = False) -> Any:
        """Search candidate resumes matching structured criteria with star ratings."""
        start_time = _time.time()
        logger.info("[search_resumes] Called with filters: %s", filters)
        if not filters or not isinstance(filters, dict):
            logger.warning("[search_resumes] Received empty or invalid filters dict")
            return [{"total_record": 0}]

        with self.get_session() as db:
            try:
                def _safe_float(value):
                    try:
                        return float(value)
                    except (TypeError, ValueError):
                        return None

                def _safe_int(value):
                    try:
                        return int(value)
                    except (TypeError, ValueError):
                        return None

                candidate_education = (
                    db.query(
                        CandidateEducation.candidate_id.label("candidate_id"),
                        func.json_agg(
                            func.json_build_object(
                                "education_id", Education.education_id,
                                "education", Education.education,
                                "institution", CandidateEducation.institution,
                                "percentage", CandidateEducation.percentage,
                                "year_of_passed", CandidateEducation.year_of_passed,
                            )
                        ).label("education"),
                        func.max(CandidateEducation.year_of_passed).label("max_year_of_passed"),
                        func.max(CandidateEducation.percentage).label("max_percentage"),
                    )
                    .select_from(CandidateEducation)
                    .join(Education, CandidateEducation.education_id == Education.education_id)
                    .group_by(CandidateEducation.candidate_id)
                    .subquery()
                )

                candidate_skills = (
                    db.query(
                        CandidateSkills.candidate_id.label("candidate_id"),
                        func.json_agg(
                            func.json_build_object(
                                "skill_id", Skill.skill_id,
                                "skill", Skill.skill,
                            )
                        ).label("skills"),
                    )
                    .select_from(CandidateSkills)
                    .join(Skill, CandidateSkills.skill_id == Skill.skill_id)
                    .group_by(CandidateSkills.candidate_id)
                    .subquery()
                )

                candidate_work_exp = (
                    db.query(
                        WorkExperience.candidate_id.label("candidate_id"),
                        func.count(WorkExperience.experience_id).label("work_exp_count"),
                        func.json_agg(
                            aggregate_order_by(
                                func.json_build_object(
                                    "role_id", Role.role_id,
                                    "role", Role.role,
                                    "company_name", Company.company_name,
                                    "company_location", Company.company_location,
                                    "start_date", WorkExperience.start_date,
                                    "end_date", WorkExperience.end_date,
                                    "is_present", WorkExperience.is_active,
                                ),
                                WorkExperience.start_date.asc(),
                            )
                        ).label("work_experience"),
                    )
                    .select_from(WorkExperience)
                    .join(Role, WorkExperience.role_id == Role.role_id)
                    .join(Company, WorkExperience.company_id == Company.company_id)
                    .group_by(WorkExperience.candidate_id)
                    .subquery()
                )

                conditions = [Candidate.is_active == True]

                name_val = filters.get("name") or filters.get("candidate_name")
                if name_val:
                    name_list = [name_val] if isinstance(name_val, str) else list(name_val)
                    name_conditions = []
                    for name_item in name_list:
                        raw_item = str(name_item).strip()
                        if not raw_item:
                            continue
                        tokens = [tok for tok in raw_item.split() if tok]
                        if len(tokens) == 1:
                            name_conditions.append(Candidate.name.ilike(f"%{tokens[0]}%"))
                        elif len(tokens) > 1:
                            name_conditions.append(and_(*[Candidate.name.ilike(f"%{tok}%") for tok in tokens]))
                    if name_conditions:
                        conditions.append(or_(*name_conditions))

                min_exp = _safe_float(filters.get("min_experience"))
                if min_exp is not None:
                    conditions.append(Candidate.total_experience >= min_exp)

                max_exp = _safe_float(filters.get("max_experience"))
                if max_exp is not None:
                    conditions.append(Candidate.total_experience <= max_exp)

                skills = filters.get("skills")
                if isinstance(skills, list) and skills:
                    conditions.append(
                        Candidate.candidate_id.in_(
                            db.query(CandidateSkills.candidate_id).filter(
                                CandidateSkills.skill_id.in_(skills)
                            )
                        )
                    )

                companies = filters.get("companies")
                if isinstance(companies, list) and companies:
                    company_terms = [str(comp).strip() for comp in companies if str(comp).strip()]
                    company_ids = []
                    if company_terms:
                        company_name_filters = [Company.company_name.ilike(f"%{term}%") for term in company_terms]
                        company_ids = [
                            row.company_id
                            for row in db.query(Company.company_id).filter(or_(*company_name_filters)).all()
                        ]

                    if company_ids:
                        conditions.append(
                            Candidate.candidate_id.in_(
                                db.query(WorkExperience.candidate_id).filter(
                                    WorkExperience.company_id.in_(company_ids)
                                )
                            )
                        )

                roles = filters.get("roles")
                if isinstance(roles, list) and roles:
                    role_terms = [str(role).strip() for role in roles if str(role).strip()]
                    role_ids = []
                    if role_terms:
                        uuid_role_ids = []
                        role_name_terms = []
                        for term in role_terms:
                            try:
                                uuid_role_ids.append(UUID(term))
                            except ValueError:
                                role_name_terms.append(term)

                        # A profile may store a generic title ("Developer") while
                        # the query contains a specialization ("Python Developer").
                        for term in list(role_name_terms):
                            title_suffix = re.search(
                                r"\b(developer|engineer|architect|analyst|designer|manager|consultant|administrator|specialist)$",
                                term, re.IGNORECASE,
                            )
                            if title_suffix:
                                role_name_terms.append(title_suffix.group(1))
                        role_name_terms = list(dict.fromkeys(role_name_terms))

                        if uuid_role_ids:
                            role_ids.extend(
                                row.role_id
                                for row in db.query(Role.role_id).filter(Role.role_id.in_(uuid_role_ids)).all()
                            )
                        if role_name_terms:
                            role_name_filters = [Role.role.ilike(f"%{term}%") for term in role_name_terms]
                            role_ids.extend(
                                row.role_id
                                for row in db.query(Role.role_id).filter(or_(*role_name_filters)).all()
                            )
                        role_ids = list(dict.fromkeys(role_ids))

                    role_conditions = []
                    if role_ids:
                        role_conditions.append(
                            Candidate.candidate_id.in_(
                                db.query(WorkExperience.candidate_id).filter(
                                    WorkExperience.role_id.in_(role_ids)
                                )
                            )
                        )
                    if role_name_terms:
                        resume_role_filters = [
                            Resume.candidate_role.ilike(f"%{term}%")
                            for term in role_name_terms
                        ]
                        role_conditions.append(
                            Candidate.candidate_id.in_(
                                db.query(Resume.candidate_id).filter(or_(*resume_role_filters))
                            )
                        )
                    if role_conditions:
                        conditions.append(or_(*role_conditions))

                education = filters.get("education")
                if isinstance(education, list) and education:
                    conditions.append(
                        Candidate.candidate_id.in_(
                            db.query(CandidateEducation.candidate_id).filter(
                                CandidateEducation.education_id.in_(education)
                            )
                        )
                    )

                year_conditions = []
                passout_start_year = _safe_int(filters.get("passout_start_year"))
                if passout_start_year is not None:
                    year_conditions.append(func.max(CandidateEducation.year_of_passed) >= passout_start_year)

                passout_end_year = _safe_int(filters.get("passout_end_year"))
                if passout_end_year is not None:
                    year_conditions.append(func.max(CandidateEducation.year_of_passed) <= passout_end_year)

                if year_conditions:
                    conditions.append(
                        Candidate.candidate_id.in_(
                            db.query(CandidateEducation.candidate_id)
                            .join(Education, Education.education_id == CandidateEducation.education_id)
                            .group_by(CandidateEducation.candidate_id)
                            .having(and_(*year_conditions))
                        )
                    )

                percentage = _safe_float(filters.get("percentage"))
                if percentage is not None:
                    conditions.append(
                        Candidate.candidate_id.in_(
                            db.query(CandidateEducation.candidate_id).filter(
                                CandidateEducation.percentage >= percentage
                            )
                        )
                    )

                total_count = (
                    db.query(func.count(Candidate.candidate_id))
                    .filter(and_(*conditions))
                    .scalar()
                ) or 0

                candidate_resumes = (
                    db.query(
                        Resume.candidate_id.label("candidate_id"),
                        func.count(Resume.resume_id).label("resume_count"),
                        func.json_agg(
                            func.json_build_object(
                                "resume_id", Resume.resume_id,
                                "candidate_role", Resume.candidate_role,
                                "file_name",
                                case(
                                    (
                                        (Resume.candidate_role.isnot(None)) & (Resume.candidate_role != ""),
                                        Resume.candidate_role,
                                    ),
                                    else_="Resume_1",
                                ),
                            )
                        ).label("resumes"),
                    )
                    .select_from(Resume)
                    .join(Attachment, Resume.attachment_id == Attachment.attachment_id)
                    .group_by(Resume.candidate_id)
                    .subquery()
                )

                recent_resume_subq = (
                    db.query(
                        Resume.candidate_id,
                        Resume.resume_id,
                        Resume.candidate_role,
                        func.row_number()
                        .over(partition_by=Resume.candidate_id, order_by=Resume.created_at.desc())
                        .label("rn"),
                    )
                    .subquery()
                )

                candidate_interviews = (
                    db.query(
                        Candidate.candidate_id.label("candidate_id"),
                        func.count(CandidateInterviews.interview_id).label("interview_count"),
                        func.json_agg(
                            func.json_build_object(
                                "interview_id", CandidateInterviews.interview_id,
                                "client_id", CandidateInterviews.client_id,
                                "resume_id", CandidateInterviews.resume_id,
                                "experience", CandidateInterviews.experience,
                                "status", CandidateInterviews.status,
                                "role", CandidateInterviews.role,
                                "created_at", CandidateInterviews.created_at,
                            )
                        ).label("interviews"),
                    )
                    .join(Resume, Resume.resume_id == CandidateInterviews.resume_id)
                    .join(Candidate, Candidate.candidate_id == Resume.candidate_id)
                    .group_by(Candidate.candidate_id)
                    .subquery()
                )

                query = (
                    db.query(
                        func.json_build_object(
                            "candidate_id", Candidate.candidate_id,
                            "name", Candidate.name,
                            "email", Candidate.email_address,
                            "phone_number", Candidate.phone_number,
                            "location", Candidate.location,
                            "total_experience", Candidate.total_experience,
                            "resume_id", recent_resume_subq.c.resume_id,
                            "candidate_role", func.coalesce(recent_resume_subq.c.candidate_role, ""),
                            "resume_count", func.coalesce(candidate_resumes.c.resume_count, 0),
                            "resumes", func.coalesce(candidate_resumes.c.resumes, cast("[]", JSON)),
                            "education", func.coalesce(candidate_education.c.education, cast("[]", JSON)),
                            "skills", func.coalesce(candidate_skills.c.skills, cast("[]", JSON)),
                            "work_experience", func.coalesce(candidate_work_exp.c.work_experience, cast("[]", JSON)),
                            "company_count", func.coalesce(candidate_work_exp.c.work_exp_count, 0),
                            "interview_count", func.coalesce(candidate_interviews.c.interview_count, 0),
                            "interviews", func.coalesce(candidate_interviews.c.interviews, cast("[]", JSON)),
                        ).label("candidate_info")
                    )
                    .select_from(Candidate)
                    .outerjoin(candidate_education, Candidate.candidate_id == candidate_education.c.candidate_id)
                    .outerjoin(candidate_skills, Candidate.candidate_id == candidate_skills.c.candidate_id)
                    .outerjoin(candidate_work_exp, Candidate.candidate_id == candidate_work_exp.c.candidate_id)
                    .outerjoin(
                        recent_resume_subq,
                        and_(Candidate.candidate_id == recent_resume_subq.c.candidate_id, recent_resume_subq.c.rn == 1),
                    )
                    .outerjoin(candidate_resumes, Candidate.candidate_id == candidate_resumes.c.candidate_id)
                    .outerjoin(candidate_interviews, Candidate.candidate_id == candidate_interviews.c.candidate_id)
                    .filter(and_(*conditions))
                )

                sort_map = {
                    "name": Candidate.name,
                    "total_experience": Candidate.total_experience,
                    "created_at": Candidate.created_at,
                    "location": Candidate.location,
                    "email": Candidate.email_address,
                    "year_of_passed": cast(candidate_education.c.max_year_of_passed, Integer),
                    "percentage": cast(candidate_education.c.max_percentage, Float),
                }

                sort_by = (filters.get("sort_by") or "name").lower()
                sort_order = (filters.get("sort_order") or "asc").lower()
                if sort_by in sort_map:
                    col = sort_map[sort_by]
                    query = query.order_by(col.desc() if sort_order == "desc" else col.asc())

                if not os.path.exists(EXPORT_PATH):
                    os.makedirs(EXPORT_PATH, exist_ok=True)

                if export:
                    export_file = os.path.join(
                        EXPORT_PATH,
                        f"exported_resumes_{datetime.datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.csv",
                    )
                    all_data = query.all()
                    if all_data:
                        df = pd.DataFrame([row.candidate_info for row in all_data])
                        df.to_csv(export_file, index=False)
                        return {"file_path": export_file}
                    return {"file_path": "No data to export"}

                page = _safe_int(filters.get("page", 1)) or 1
                page_size = _safe_int(filters.get("page_size", 20)) or 20
                page = max(1, page)
                page_size = max(1, min(100, page_size))
                query = query.limit(page_size).offset((page - 1) * page_size)

                data = query.all()
                jd_role_values = filters.get("roles") or []
                jd_role_ids = []
                jd_role_names = []
                for value in jd_role_values:
                    try:
                        jd_role_ids.append(UUID(str(value)))
                    except (ValueError, TypeError):
                        if str(value).strip():
                            jd_role_names.append(str(value).strip())
                if jd_role_ids:
                    db_roles = [r.role for r in db.query(Role.role).filter(Role.role_id.in_(jd_role_ids)).all()]
                else:
                    db_roles = []
                jd_roles = list(dict.fromkeys(db_roles + jd_role_names))

                jd_skill_values = filters.get("skills") or []
                jd_skill_ids = []
                jd_skill_names = []
                for item in jd_skill_values:
                    try:
                        jd_skill_ids.append(UUID(str(item)))
                    except (ValueError, TypeError):
                        if str(item).strip():
                            jd_skill_names.append(str(item).strip())
                if jd_skill_ids:
                    db_skills = [
                        skill.skill for skill in db.query(Skill.skill).filter(Skill.skill_id.in_(jd_skill_ids)).all()
                    ]
                else:
                    db_skills = []
                jd_skills = list(dict.fromkeys(db_skills + jd_skill_names))

                results = []
                for row in data:
                    candidate_info = row.candidate_info
                    candidate_skills_list = candidate_info.get("skills", [])
                    rate_skill = [
                        item.get("skill") if isinstance(item, dict) else str(item)
                        for item in candidate_skills_list
                        if (isinstance(item, dict) and item.get("skill")) or (isinstance(item, str) and item.strip())
                    ]

                    candidate_roles = []
                    resume_role = candidate_info.get("candidate_role", "") or ""
                    if resume_role.strip():
                        candidate_roles.append(resume_role.strip())
                    for exp in candidate_info.get("work_experience", []):
                        if isinstance(exp, dict) and exp.get("role"):
                            candidate_roles.append(str(exp["role"]).strip())

                    candidate_rating = calculate_match_score(
                        jd_role=jd_roles,
                        candidate_role=candidate_roles if candidate_roles else resume_role,
                        jd_skills=jd_skills,
                        candidate_skills=rate_skill,
                    )
                    candidate_info["candidate_rating"] = candidate_rating.get("star_rating")
                    results.append(candidate_info)

                results.append({"total_record": total_count})
                logger.info("[search_resumes] Completed in %ss", round(_time.time() - start_time, 3))
                return results

            except Exception as e:
                logger.error("[search_resumes] Unexpected error: %s", str(e), exc_info=True)
                raise

    def execute_search(
        self,
        resolved_filters: Optional[dict],
        page: int,
        page_size: int,
        sort_by: Optional[str],
        sort_order: str,
        export: bool = False,
    ):
        """Validate search inputs, inject pagination and invoke search_resumes."""
        try:
            if resolved_filters is None:
                resolved_filters = {}

            if not isinstance(resolved_filters, dict):
                raise ValueError("resolved_filters must be a dictionary")

            filter_dict = dict(resolved_filters)

            if not export:
                if page is None or page_size is None:
                    raise ValueError("page and page_size are required")
                if not isinstance(page, int) or not isinstance(page_size, int):
                    raise ValueError("page and page_size must be integers")
                if page < 1 or page_size < 1:
                    raise ValueError("page and page_size must be greater than 0")
                if page_size > 100:
                    raise ValueError("page_size cannot exceed 100")

                if sort_order:
                    sort_order = sort_order.lower().strip()
                    if sort_order not in ["asc", "desc"]:
                        raise ValueError("sort_order must be either 'asc' or 'desc'")

                filter_dict["page"] = page
                filter_dict["page_size"] = page_size
                if sort_by:
                    filter_dict["sort_by"] = sort_by
                filter_dict["sort_order"] = sort_order

            try:
                parsed = ResumeFilterRequest(**filter_dict).model_dump()
            except ValidationError:
                parsed = filter_dict

            if not export:
                parsed["page"] = page
                parsed["page_size"] = page_size

            rows = self.search_resumes(parsed, export)

            if export:
                csv_file_path = rows.get("file_path") if isinstance(rows, dict) else None
                if not csv_file_path:
                    return {
                        "status": status.HTTP_500_INTERNAL_SERVER_ERROR,
                        "message": "Export file path not found",
                    }
                if not os.path.exists(csv_file_path):
                    return {
                        "status": status.HTTP_500_INTERNAL_SERVER_ERROR,
                        "message": "Export file does not exist",
                    }
                with open(csv_file_path, "rb") as file:
                    base64_string = base64.b64encode(file.read()).decode("utf-8")
                return {
                    "status": status.HTTP_200_OK,
                    "message": "Export completed successfully",
                    "data": {"content": base64_string},
                }

            total_record = 0
            candidates = []
            if rows:
                for item in rows:
                    if isinstance(item, dict) and "total_record" in item:
                        total_record = item["total_record"]
                    else:
                        candidates.append(item)

            return candidates, total_record

        except Exception as e:
            raise Exception(f"Error while executing resume search: {str(e)}")

    def export_candidate(self, search_id: str, page: int, page_size: int, export: bool):
        """Export candidates matching a cached search session as base64-encoded CSV."""
        try:
            result = self.load_search_session(search_id)
            if not result:
                raise CandidateServiceError(f"Search session '{search_id}' not found", status_code=404)
            resolved_filters = result["filters"]
            data = self.execute_search(
                resolved_filters,
                page,
                page_size,
                sort_by=None,
                sort_order="asc",
                export=export,
            )
            return serialize_response(data)
        except CandidateServiceError as e:
            raise HTTPException(
                status_code=e.status_code,
                detail={"status": "error", "message": e.message},
            )
        except HTTPException:
            raise
        except Exception as e:
            logger.warning("[export_candidate] Error: %s", str(e), exc_info=True)
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail={"status": "error", "message": str(e)},
            )

    def nl_search_initial(
        self,
        user_query: str,
        page: Optional[int] = 1,
        page_size: Optional[int] = 20,
        sort_by: Optional[str] = None,
        sort_order: Optional[str] = "asc",
        export: bool = False,
    ) -> dict:
        """First-time NL search: LLM -> resolve -> cache -> query -> respond."""
        _page = page or 1
        _page_size = page_size or 20
        _sort_order = sort_order or "asc"

        raw_filters, resolved_filters = self.extract_and_resolve(user_query)

        search_id = uuid4().hex
        self.store_search_session(search_id, raw_filters, resolved_filters, user_query)

        candidates, total_record = self.execute_search(
            resolved_filters, _page, _page_size, sort_by, _sort_order, export
        )

        return {
            "status": "success",
            "search_id": search_id,
            "user_query": user_query,
            "filters_applied": raw_filters,
            "candidates": candidates,
            "total_record": total_record,
            "page": _page,
            "page_size": _page_size,
        }

    def nl_search_paginate(
        self,
        search_id: str,
        page: Optional[int] = 1,
        page_size: Optional[int] = 20,
        sort_by: Optional[str] = None,
        sort_order: Optional[str] = "asc",
        export: bool = False,
    ) -> dict:
        """Follow-up request: load cached filters -> query -> respond."""
        _page = page or 1
        _page_size = page_size or 20
        _sort_order = sort_order or "asc"

        session = self.load_search_session(search_id)
        if session is None:
            raise LookupError(f"search_id '{search_id}' not found or expired. Please start a new search.")

        resolved_filters = session["filters"]
        raw_filters = session.get("raw_filters", resolved_filters)
        user_query = session.get("user_query")

        candidates, total_record = self.execute_search(
            resolved_filters, _page, _page_size, sort_by, _sort_order, export
        )

        return {
            "status": "success",
            "search_id": search_id,
            "user_query": user_query,
            "filters_applied": raw_filters,
            "candidates": candidates,
            "total_record": total_record,
            "page": _page,
            "page_size": _page_size,
        }

    def get_master_data(self) -> dict:
        """Fetch active master data: Roles, Education, Skills."""
        with self.get_session() as db:
            try:
                roles = db.query(Role).filter(Role.is_active == True).all()
                educations = db.query(Education).filter(Education.is_active == True).all()
                skills = db.query(Skill).filter(Skill.is_active == True).all()

                return {
                    "roles": [{"id": str(r.role_id), "name": r.role} for r in roles],
                    "education": [{"id": str(e.education_id), "name": e.education} for e in educations],
                    "skills": [{"id": str(s.skill_id), "name": s.skill} for s in skills],
                }
            except Exception as e:
                logger.error("[get_master_data] Error fetching master data: %s", str(e), exc_info=True)
                raise

    def apply_filters(self, candidates: list, filters: dict) -> list:
        """Apply filters in-memory on candidate records returned from DB."""
        filtered_candidates = []
        filter_dict = filters.get("filters") if (isinstance(filters, dict) and "filters" in filters) else filters
        if not isinstance(filter_dict, dict):
            filter_dict = {}

        name_filter_val = filter_dict.get("name") or filter_dict.get("candidate_name")
        name_entries = []
        if name_filter_val:
            if isinstance(name_filter_val, str):
                name_entries = [name_filter_val.strip()]
            elif isinstance(name_filter_val, list):
                name_entries = [str(n).strip() for n in name_filter_val if str(n).strip()]

        for candidate in candidates:
            matched = True

            if name_entries:
                candidate_name = (candidate.get("name") or "").lower()

                def _matches_name_entry(entry: str, cand_name: str) -> bool:
                    tokens = [t.lower() for t in entry.split() if t]
                    if not tokens:
                        return True
                    return all(tok in cand_name for tok in tokens)

                if not any(_matches_name_entry(entry, candidate_name) for entry in name_entries):
                    matched = False

            if matched and filter_dict.get("location"):
                location_filter = filter_dict["location"].lower()
                candidate_location = (candidate.get("location") or "").lower()
                if location_filter not in candidate_location:
                    matched = False

            if matched and filter_dict.get("skills"):
                candidate_skill_ids = set()
                for skill in candidate.get("skills", []):
                    if isinstance(skill, dict) and skill.get("skill_id") is not None:
                        candidate_skill_ids.add(skill.get("skill_id"))

                filter_skill_ids = set(filter_dict["skills"])
                if not filter_skill_ids.issubset(candidate_skill_ids):
                    matched = False

            if matched and filter_dict.get("companies"):
                candidate_companies = set()
                for exp in candidate.get("work_experience", []):
                    if isinstance(exp, dict):
                        company_name = exp.get("company_name", "")
                    elif isinstance(exp, str):
                        company_name = exp
                    else:
                        continue
                    if company_name:
                        candidate_companies.add(company_name.lower())

                filter_companies = {str(company).lower() for company in filter_dict["companies"]}
                if not any(company in candidate_companies for company in filter_companies):
                    matched = False

            if matched and filter_dict.get("passout_start_year"):
                start_year = int(filter_dict["passout_start_year"])
                education_years = []
                for edu in candidate.get("education", []):
                    if isinstance(edu, dict) and edu.get("year_of_passed") is not None:
                        try:
                            education_years.append(int(edu.get("year_of_passed")))
                        except (ValueError, TypeError):
                            pass
                if not any(year >= start_year for year in education_years):
                    matched = False

            if matched and filter_dict.get("passout_end_year"):
                end_year = int(filter_dict["passout_end_year"])
                education_years = []
                for edu in candidate.get("education", []):
                    if isinstance(edu, dict) and edu.get("year_of_passed") is not None:
                        try:
                            education_years.append(int(edu.get("year_of_passed")))
                        except (ValueError, TypeError):
                            pass
                if not any(year <= end_year for year in education_years):
                    matched = False

            if matched and filter_dict.get("percentage"):
                required_percentage = float(filter_dict["percentage"])
                percentages = []
                for edu in candidate.get("education", []):
                    if isinstance(edu, dict) and edu.get("percentage") is not None:
                        try:
                            percentages.append(float(edu.get("percentage")))
                        except (ValueError, TypeError):
                            pass
                if not any(p >= required_percentage for p in percentages):
                    matched = False

            if matched and filter_dict.get("experience"):
                filter_experience = {int(exp) for exp in filter_dict["experience"]}
                try:
                    candidate_experience = int(candidate.get("total_experience"))
                except (ValueError, TypeError):
                    matched = False

                if matched and candidate_experience not in filter_experience:
                    matched = False

            if matched:
                filtered_candidates.append(candidate)

        return filtered_candidates

    def get_candidates(
        self,
        filters: Optional[Dict[str, Any]] = None,
        page: int = 1,
        page_size: int = 20,
    ) -> List[Dict[str, Any]]:
        """Retrieve candidate information with optional filtering, sorting, and pagination."""
        try:
            if not filters:
                return self.get_candidate_details(page, None, None)

            try:
                parsed_filters = ResumeFilterRequest(**filters).model_dump()
            except ValidationError as ve:
                raise HTTPException(
                    status_code=400,
                    detail={
                        "status": "error",
                        "message": "Invalid filter parameters",
                        "errors": [
                            {
                                "field": ".".join(str(loc) for loc in err.get("loc", [])),
                                "message": err.get("msg", ""),
                            }
                            for err in ve.errors()
                        ],
                    },
                )

            parsed_filters["page"] = page
            parsed_filters["page_size"] = page_size

            rows = self.search_resumes(parsed_filters)
            return serialize_response(rows)

        except CandidateServiceError as e:
            raise HTTPException(
                status_code=e.status_code,
                detail={"status": "error", "message": e.message},
            )
        except HTTPException:
            raise
        except Exception as e:
            logger.error("Error in get_candidates: %s", str(e), exc_info=True)
            raise HTTPException(
                status_code=500,
                detail={"status": "error", "message": "Failed to fetch candidates"},
            )

    def generate_dynamic_filters(self, body: dict) -> Dict[str, Any]:
        """Generate structured filter payload from a natural-language recruiter query."""
        start_time = _time.time()
        logger.info("[generate_dynamic_filters] Request received | body_keys=%s", list(body.keys()) if body else "empty")

        query = (body or {}).get("query", "")
        if not query or not isinstance(query, str) or not query.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail={"status": "error", "message": "A non-empty 'query' string is required"},
            )

        try:
            req = DynamicFilterRequest(query=query)
        except ValidationError as ve:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail={
                    "status": "error",
                    "message": "Invalid request",
                    "errors": [
                        {"field": ".".join(str(loc) for loc in err.get("loc", [])), "message": err.get("msg", "")}
                        for err in ve.errors()
                    ],
                },
            )

        try:
            raw_filters = self.extract_filters_from_query(req.query)
            if not raw_filters:
                raw_filters, _ = self._fallback_extract_and_resolve(req.query)
            logger.info("[generate_dynamic_filters] Extracted filters: %s", raw_filters)

            try:
                validated = DynamicFilterResponse(**raw_filters).model_dump(exclude_none=True)
            except ValidationError:
                validated = raw_filters

            elapsed = round(_time.time() - start_time, 3)
            logger.info("[generate_dynamic_filters] Completed in %ss", elapsed)
            return {"status": "success", "data": validated}

        except HTTPException:
            raise
        except Exception as e:
            elapsed = round(_time.time() - start_time, 3)
            logger.error("[generate_dynamic_filters] Error after %ss: %s", elapsed, str(e), exc_info=True)
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail={"status": "error", "message": "Failed to generate dynamic filters"},
            )

    def filter_resumes(
        self,
        filters: dict,
        search_id: str,
        sort_by: Optional[str] = None,
        sort_order: Optional[str] = None,
        page: int = 1,
        page_size: int = 20,
    ) -> dict:
        """Filter resumes stored in the database with pagination support."""
        start_time = _time.time()
        logger.info(
            "[filter_resumes] Request received | page=%s, page_size=%s, payload_keys=%s",
            page, page_size, list(filters.keys()) if filters else "empty"
        )

        if not filters or not isinstance(filters, dict):
            logger.warning("[filter_resumes] Empty or invalid request body received")
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Request body must be a non-empty JSON object",
            )

        try:
            is_combined = "filters" in filters or "dynamic_filters" in filters

            if is_combined:
                logger.info("[filter_resumes] Combined filter mode detected (standard + dynamic)")
                standard_raw = filters.get("filters") or {}
                dynamic_raw = filters.get("dynamic_filters") or {}

                resolved_dynamic = self.resolve_dynamic_filters(dynamic_raw) if dynamic_raw else {}
                logger.info("[filter_resumes] Resolved dynamic filters: %s", resolved_dynamic)

                merged = self.merge_filters(standard_raw, resolved_dynamic)

                try:
                    parsed_filters = ResumeFilterRequest(**merged).model_dump()
                except ValidationError as ve:
                    logger.warning("[filter_resumes] Merged filter validation failed: %s", ve.errors())
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail={
                            "status": "error",
                            "message": "Invalid filter parameters",
                            "errors": [
                                {"field": ".".join(str(loc) for loc in err.get("loc", [])), "message": err.get("msg", "")}
                                for err in ve.errors()
                            ],
                        },
                    )

                parsed_filters["page"] = page
                parsed_filters["page_size"] = page_size

                rows = self.search_resumes(parsed_filters)
                result_count = len(rows) - 1 if rows else 0
                logger.info("[filter_resumes] Combined search returned %s results", result_count)
                elapsed = round(_time.time() - start_time, 3)
                logger.info("[filter_resumes] Request completed in %ss", elapsed)
                return serialize_response(rows)

            else:
                logger.info("[filter_resumes] Structured filter mode detected")
                try:
                    parsed_filters = ResumeFilterRequest(**filters).model_dump()
                except ValidationError as ve:
                    logger.warning("[filter_resumes] Filter validation failed: %s", ve.errors())
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail={
                            "status": "error",
                            "message": "Invalid filter parameters",
                            "errors": [
                                {"field": ".".join(str(loc) for loc in err.get("loc", [])), "message": err.get("msg", "")}
                                for err in ve.errors()
                            ],
                        },
                    )

                parsed_filters["page"] = page
                parsed_filters["page_size"] = page_size

                session = self.load_search_session(search_id)
                if session is None:
                    raise LookupError(
                        f"search_id '{search_id}' not found or expired. "
                        "Please start a new search."
                    )

                resolved_filters = session["filters"]
                candidates, total_record = self.execute_search(
                    resolved_filters, page, page_size, sort_by, sort_order, export=False
                )
                result = self.apply_filters(candidates, filters)

                return {
                    "status": status.HTTP_200_OK,
                    "message": "filter records successfully",
                    "data": {
                        "candidates": result,
                        "total_records": len(result),
                        "search_id": search_id,
                    },
                }

        except HTTPException:
            raise

        except LookupError as le:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail={"status": "error", "message": str(le)},
            )

        except Exception as e:
            elapsed = round(_time.time() - start_time, 3)
            logger.error("[filter_resumes] Unexpected error after %ss: %s", elapsed, str(e), exc_info=True)
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail={
                    "status": "error",
                    "message": "Failed to filter resumes",
                },
            )

    def profile_search(self, body: dict) -> Dict[str, Any]:
        """Search candidates using a natural-language query with query caching and pagination."""
        start_time = _time.time()
        logger.info("[profile_search] Request received | keys=%s", list(body.keys()) if body else "empty")

        if not body or not isinstance(body, dict):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail={"status": "error", "message": "Request body must be a non-empty JSON object"},
            )

        search_id = body.get("search_id")

        try:
            if search_id:
                req = NLSearchPaginateRequest(**body)
                result = self.nl_search_paginate(
                    search_id=req.search_id,
                    page=req.page,
                    page_size=req.page_size,
                    sort_by=req.sort_by,
                    sort_order=req.sort_order,
                )
            else:
                if not body.get("user_query"):
                    raise HTTPException(
                        status_code=status.HTTP_400_BAD_REQUEST,
                        detail={
                            "status": "error",
                            "message": "Either 'user_query' or 'search_id' is required",
                        },
                    )
                req = NLSearchRequest(**body)
                result = self.nl_search_initial(
                    user_query=req.user_query,
                    page=req.page,
                    page_size=req.page_size,
                    sort_by=req.sort_by,
                    sort_order=req.sort_order,
                )

            result["candidates"] = serialize_response(result.get("candidates", []))
            elapsed = round(_time.time() - start_time, 3)
            logger.info("[profile_search] Completed in %ss | search_id=%s", elapsed, result.get("search_id"))
            return result

        except ValidationError as ve:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail={
                    "status": "error",
                    "message": "Invalid request parameters",
                    "errors": [
                        {"field": ".".join(str(loc) for loc in err.get("loc", [])), "message": err.get("msg", "")}
                        for err in ve.errors()
                    ],
                },
            )

        except LookupError as le:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail={"status": "error", "message": str(le)},
            )

        except ValueError as ve:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail={"status": "error", "message": str(ve)},
            )

        except HTTPException:
            raise

        except Exception as e:
            elapsed = round(_time.time() - start_time, 3)
            logger.error("[profile_search] Unexpected error after %ss: %s", elapsed, str(e), exc_info=True)
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail={"status": "error", "message": "Natural language search failed"},
            )

    def get_filter_options(self) -> Dict[str, Any]:
        """Fetch master data for filters: Roles, Education, Skills."""
        try:
            data = self.get_master_data()
            return {"status": "success", "data": data}
        except Exception as e:
            logger.error("ERROR in get_filter_options: %s", str(e), exc_info=True)
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail={
                    "status": "error",
                    "message": "Failed to fetch master data",
                },
            )


# =============================================================================
# 4. Candidate Parser Service
# =============================================================================

class CandidateParserService(BaseService):
    """Encapsulates document text extraction, resume classification, information extraction, and DB storage."""

    @staticmethod
    def extract_text_from_pdf(path: str) -> str:
        """Extract plain text from PDF pages."""
        text = ""
        try:
            reader = PdfReader(path)
            for page in reader.pages:
                text += page.extract_text() or ""
            logger.info("Extracted text from PDF: %s", path)
        except Exception as e:
            logger.info("[ERROR] PDF read failed: %s -> %s", path, e)
        return text

    @staticmethod
    def extract_docx_text(path: str) -> str:
        """Extract text from docx paragraphs, tables, and XML elements."""
        full_text = []
        try:
            doc = Document(path)
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    full_text.append(text)

            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            row_text.append(cell_text)
                    if row_text:
                        full_text.append(" | ".join(row_text))

            with zipfile.ZipFile(path) as z:
                xml_content = z.read("word/document.xml")

            tree = etree.XML(xml_content)
            namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            xml_texts = tree.xpath("//w:t/text()", namespaces=namespaces)

            for text in xml_texts:
                cleaned = text.strip()
                if cleaned and cleaned not in full_text:
                    full_text.append(cleaned)

            unique_text = list(dict.fromkeys(full_text))
            return "\n".join(unique_text)
        except Exception as e:
            logger.error("[ERROR] DOCX extraction failed: %s -> %s", path, e)
            return ""

    @staticmethod
    def extract_text_from_docx(path: str) -> str:
        """Simple paragraph-only docx extraction."""
        text = ""
        try:
            doc = Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
        except Exception as e:
            logger.info("[ERROR] DOCX read failed: %s -> %s", path, e)
        return text

    def run_llm_json(
        self, prompt: str, text: str, system: str = "You only return JSON", max_retries: int = 3
    ) -> dict:
        """Select available model and call it with retries."""
        with self.get_session() as db:
            # Find an admin that has at least one active model config
            admin = (
                db.query(Admin)
                .join(AiModelConfig, AiModelConfig.admin_id == Admin.admin_id)
                .filter(Admin.is_active == True, AiModelConfig.is_active == True)
                .first()
            )
            if not admin:
                # Fallback: any active admin
                admin = db.query(Admin).filter(Admin.is_active == True).first()
            if not admin:
                raise CandidateServiceError("No admin found", status_code=500)
            admin_id = admin.admin_id
            admin_name = admin.name

        failed_ids = set()
        last_error = None

        for _ in range(max_retries):
            model_info = select_available_model(admin_id, exclude_ids=failed_ids)
            if not model_info:
                break

            model_config_id = model_info.get("model_config_id")
            model = model_info.get("model_name", "ollama").lower()
            version = model_info.get("version_name", "llama3").lower()
            base_url = model_info.get("base_url")
            api_key = decrypt_data(model_info.get("apikey"))

            message = [
                {"role": "system", "content": system},
                {"role": "user", "content": prompt + "\n\nDocument:\n" + text[:12000]},
            ]

            try:
                if model == "openai":
                    if not api_key:
                        raise CandidateServiceError("OpenAI API key not found", status_code=500)
                    client = OpenAI(
                        api_key=api_key,
                        base_url=base_url or "https://api.openai.com/v1",
                    )
                    response = client.chat.completions.create(model=version, messages=message)
                    content = response.choices[0].message.content.strip()
                    record_model_usage(model_config_id, response.usage.total_tokens, admin_name)

                elif model == "groq":
                    if not api_key:
                        raise CandidateServiceError("Groq API key not found", status_code=500)
                    client = OpenAI(
                        api_key=api_key,
                        base_url=base_url or "https://api.groq.com/openai/v1",
                    )
                    response = client.chat.completions.create(
                        model=version,
                        messages=message,
                        response_format={"type": "json_object"},
                    )
                    content = response.choices[0].message.content.strip()
                    if response.usage:
                        record_model_usage(model_config_id, response.usage.total_tokens, admin_name)

                elif model == "claude":
                    if not api_key:
                        raise CandidateServiceError("Claude API key not found", status_code=500)
                    client = Anthropic(
                        api_key=api_key,
                        base_url=base_url or "https://api.anthropic.com",
                    )
                    response = client.messages.create(
                        model=version,
                        max_tokens=1000,
                        system=system,
                        messages=[{"role": "user", "content": prompt + "\n\nDocument:\n" + text[:12000]}],
                    )
                    content = response.content[0].text.strip()
                    token_used = response.usage.input_tokens + response.usage.output_tokens
                    record_model_usage(model_config_id, token_used, admin_name)

                else:
                    response = ollama.chat(model=version, messages=message)
                    content = response["message"]["content"].strip()

                start = content.find("{")
                end = content.rfind("}") + 1
                if start == -1 or end == 0:
                    raise ValueError("No valid JSON object found in model response")

                return json.loads(content[start:end])

            except Exception as call_err:
                last_error = call_err
                logger.warning("[run_llm_json] model=%s failed, trying next: %s", model, call_err)
                failed_ids.add(model_config_id)
                continue

        raise CandidateServiceError(f"All AI model attempts failed: {last_error}", status_code=500)

    @staticmethod
    def is_resume(text: str) -> bool:
        """Classify whether a document is a resume/CV."""
        prompt = """ 
        You are a strict classifier.

        Return ONLY JSON:
        {"is_resume": true} or {"is_resume": false}

        Rules:
        - Return TRUE only if this is clearly a complete resume/CV
        - A resume MUST contain at least 2 of these sections: Skills, Experience, Education, Projects
        - Return FALSE if it is invoice, email, report, or random text
        """
        data = llm_service.call_llm(
            prompt=prompt + "\n\nDocument:\n" + text[:12000],
            system_prompt="You only return JSON",
            is_json=True,
        )
        return bool(data.get("is_resume", False))

    @staticmethod
    def extract_basic_info(resume_text: str) -> dict:
        """Extract structured candidate information from resume text using LLM."""
        prompt = """
        You are a highly accurate resume parser.
        Extract structured candidate information from the given resume.
        Return ONLY valid JSON.

        STRICT JSON FORMAT:
        {
            "name": "",
            "total_experience": 0,
            "email": "",
            "phone_number": "",
            "location": "",
            "role": "",
            "skills": [],
            "education": [{"qualification": "", "institution": "", "percentage": "", "passout_year": ""}],
            "work_experience": [{"company_name": "", "role": "", "start_date": "", "end_date": ""}]
        }

        Rules:
        - For work_experience: Extract all work experience, company employments, client engagements, or projects worked on. If an explicit employer/company is given, use it. If client organizations or projects are listed without a distinct employer header (e.g. 'Hollweg Assessment Partners', 'Wagepoint'), use the client or project name as company_name. If no company or client is named, use 'Independent'.
        - For role in work_experience: extract the role or job title for that experience (e.g. 'Software Engineer', 'SQL Developer'). If not explicitly stated per job, infer it from the candidate's primary role.
        - For start_date and end_date in work_experience: extract in "YYYY-MM" or "YYYY" format (or "present" for current role). If start_date or end_date is NOT provided in the resume, leave it as null or empty string. Do NOT invent dates or use the current date.
        """
        data = llm_service.call_llm(
            prompt=prompt + "\n\nDocument:\n" + resume_text[:12000],
            system_prompt="You only return JSON",
            is_json=True,
            # max_tokens=2000,
        )
        try:
            logger.info(
                "[extract_basic_info] LLM successfully extracted JSON:\n%s",
                json.dumps(data, indent=2, default=str) if isinstance(data, (dict, list)) else data,
            )
        except Exception:
            logger.info("[extract_basic_info] LLM successfully extracted JSON: %s", data)

        return data

    @staticmethod
    def clean_json_response(content: Union[str, dict]) -> dict:
        """Extract valid JSON from LLM response."""
        if isinstance(content, dict):
            return content

        match = re.search(r"\{.*\}", content, re.DOTALL)
        if not match:
            raise ValueError("No valid JSON found in response")

        return json.loads(match.group(0))

    def save_resumes_to_db(self, resumes: List[dict]) -> None:
        """Save extracted resume info to PostgreSQL with normalized schema."""
        with self.get_session() as db:
            for r in resumes:
                try:
                    info = r["info"]
                    attachment_id = r.get("attachment_id")
                    sender_email = parse_email_address(r.get("sender_email", ""))

                    info_email = (info.get("email") or "").strip().lower()
                    extracted_role = info.get("role", "").strip()
                    if not info_email and sender_email:
                        info_email = sender_email

                    email_address = info_email or sender_email or ""
                    phone_number = clean_mobile_number((info.get("phone_number") or "").strip())

                    candidate = None
                    if email_address:
                        candidate = (
                            db.query(Candidate)
                            .filter(func.lower(Candidate.email_address) == email_address.lower())
                            .first()
                        )

                    if not candidate and phone_number:
                        candidate = db.query(Candidate).filter(Candidate.phone_number == phone_number).first()

                    if candidate:
                        if not candidate.email_address and email_address:
                            candidate.email_address = email_address
                            candidate.email_from_sender = bool(sender_email and not info_email)
                        if not candidate.phone_number and phone_number:
                            candidate.phone_number = phone_number
                        db.add(candidate)
                    else:
                        candidate = Candidate(
                            name=info.get("name", ""),
                            email_address=email_address,
                            email_from_sender=bool(sender_email and not info_email),
                            phone_number=phone_number,
                            location=info.get("location", ""),
                            total_experience=normalize_experience(info.get("total_experience")),
                            created_by="resume_parser",
                        )
                        db.add(candidate)
                        db.flush()

                    # Skills
                    skills_list = info.get("skills", [])
                    if skills_list:
                        for skill_name in skills_list:
                            if skill_name and isinstance(skill_name, str):
                                skill_name = skill_name.strip()
                                skill = db.query(Skill).filter(Skill.skill.ilike(skill_name)).first()
                                if not skill:
                                    skill = Skill(skill=skill_name, created_by="resume_parser")
                                    db.add(skill)
                                    db.flush()

                                is_skill = (
                                    db.query(CandidateSkills)
                                    .filter(
                                        CandidateSkills.candidate_id == candidate.candidate_id,
                                        CandidateSkills.skill_id == skill.skill_id,
                                    )
                                    .first()
                                )
                                if not is_skill:
                                    candidate_skill = CandidateSkills(
                                        candidate_id=candidate.candidate_id,
                                        skill_id=skill.skill_id,
                                        created_by="resume_parser",
                                    )
                                    db.add(candidate_skill)

                    # Education
                    education_list = info.get("education", [])
                    if education_list:
                        for edu_item in education_list:
                            if isinstance(edu_item, dict):
                                qualification = edu_item.get("qualification", "").strip()
                                institution = edu_item.get("institution", "").strip()
                                percentage = edu_item.get("percentage", "")
                                try:
                                    value = float(str(percentage).replace("%", "").strip())
                                    if value <= 10:
                                        percentage = str(round(value * 9.5, 2))
                                    else:
                                        percentage = str(value)
                                except (ValueError, AttributeError):
                                    percentage = ""

                                passout_year = edu_item.get("passout_year", "")
                                if qualification:
                                    education_type = (
                                        db.query(Education)
                                        .filter(Education.education.ilike(qualification))
                                        .first()
                                    )
                                    if not education_type:
                                        education_type = Education(
                                            education=qualification, created_by="resume_parser"
                                        )
                                        db.add(education_type)
                                        db.flush()

                                    try:
                                        year_passed = int(passout_year) if passout_year else None
                                    except (ValueError, TypeError):
                                        year_passed = None

                                    try:
                                        percentage_val = float(percentage) if percentage else None
                                    except (ValueError, TypeError):
                                        percentage_val = None

                                    is_education = (
                                        db.query(CandidateEducation)
                                        .filter(
                                            CandidateEducation.candidate_id == candidate.candidate_id,
                                            CandidateEducation.education_id == education_type.education_id,
                                        )
                                        .first()
                                    )
                                    if not is_education:
                                        candidate_edu = CandidateEducation(
                                            candidate_id=candidate.candidate_id,
                                            education_id=education_type.education_id,
                                            institution=institution or None,
                                            percentage=percentage_val,
                                            year_of_passed=year_passed,
                                            created_by="resume_parser",
                                        )
                                        db.add(candidate_edu)

                    # Work Experience
                    work_experience_list = info.get("work_experience", [])
                    if work_experience_list:
                        for work_exp_item in work_experience_list:
                            if isinstance(work_exp_item, dict):
                                company_name = work_exp_item.get("company_name", "").strip()
                                role_name = work_exp_item.get("role", "").strip() or extracted_role or "Software Professional"
                                start_date = work_exp_item.get("start_date", "")
                                end_date = work_exp_item.get("end_date", "")

                                if not company_name and (role_name or start_date or end_date):
                                    company_name = "Independent"

                                if company_name:
                                    company = (
                                        db.query(Company)
                                        .filter(Company.company_name.ilike(company_name))
                                        .first()
                                    )
                                    if not company:
                                        company = Company(
                                            company_name=company_name,
                                            company_location="",
                                            created_by="resume_parser",
                                        )
                                        db.add(company)
                                        db.flush()

                                    role_obj = db.query(Role).filter(Role.role.ilike(role_name)).first()
                                    if not role_obj:
                                        role_obj = Role(role=role_name, created_by="resume_parser")
                                        db.add(role_obj)
                                        db.flush()

                                    start_dt = None
                                    end_dt = None
                                    try:
                                        if start_date:
                                            s_str = str(start_date).strip()
                                            if s_str and s_str.lower() not in {"null", "none", "n/a", "na", "undefined", "present", ""}:
                                                if len(s_str) == 4 and s_str.isdigit():
                                                    start_dt = datetime.date(int(s_str), 1, 1)
                                                elif len(s_str) == 7 and "-" in s_str:
                                                    parts = s_str.split("-")
                                                    if len(parts) == 2 and parts[0].isdigit() and parts[1].isdigit():
                                                        start_dt = datetime.date(int(parts[0]), int(parts[1]), 1)
                                    except (ValueError, TypeError):
                                        start_dt = None

                                    try:
                                        if end_date:
                                            e_str = str(end_date).strip()
                                            if e_str and e_str.lower() not in {"null", "none", "n/a", "na", "undefined", "present", ""}:
                                                if len(e_str) == 4 and e_str.isdigit():
                                                    end_dt = datetime.date(int(e_str), 12, 31)
                                                elif len(e_str) == 7 and "-" in e_str:
                                                    parts = e_str.split("-")
                                                    if len(parts) == 2 and parts[0].isdigit() and parts[1].isdigit():
                                                        year, month = parts
                                                        if month == "12":
                                                            end_dt = datetime.date(int(year), 12, 31)
                                                        else:
                                                            next_month = datetime.date(int(year), int(month) + 1, 1)
                                                            end_dt = next_month - datetime.timedelta(days=1)
                                    except (ValueError, TypeError):
                                        end_dt = None

                                    is_work_exp = (
                                        db.query(WorkExperience)
                                        .filter(
                                            WorkExperience.candidate_id == candidate.candidate_id,
                                            WorkExperience.company_id == company.company_id,
                                            WorkExperience.role_id == role_obj.role_id,
                                        )
                                        .first()
                                    )
                                    if not is_work_exp:
                                        work_exp = WorkExperience(
                                            candidate_id=candidate.candidate_id,
                                            company_id=company.company_id,
                                            role_id=role_obj.role_id,
                                            start_date=start_dt,
                                            end_date=end_dt,
                                            created_by="resume_parser",
                                        )
                                        db.add(work_exp)

                    # Resume
                    is_resume_record = (
                        db.query(Resume)
                        .filter(
                            Resume.candidate_id == candidate.candidate_id,
                            Resume.candidate_role == extracted_role,
                        )
                        .first()
                    )
                    if not is_resume_record:
                        resume_record = Resume(
                            candidate_id=candidate.candidate_id,
                            attachment_id=attachment_id,
                            candidate_role=extracted_role,
                            created_by="resume_parser",
                        )
                        db.add(resume_record)
                        db.flush()

                    db.commit()

                except Exception as e:
                    logger.error("Error saving resume to DB: %s", e, exc_info=True)
                    db.rollback()
                    continue

    def process_resumes(self, email_id: UUID) -> dict:
        """Classify and process resume attachments for a given email (invoked by Celery)."""
        logger.info("process_resumes called for email_id=%s", email_id)
        with self.get_session() as db:
            email_obj = db.query(EmailLogs).filter(EmailLogs.email_id == email_id).first()
            if not email_obj:
                return {"message": "Email not found"}

            all_attachments = (
                db.query(Attachment).filter(Attachment.email_id == email_obj.email_id).all()
            )

            results = []
            kept_files = []

            for att in all_attachments:
                try:
                    file_path = attachment_storage_service.stage_attachment(att.file_name)
                    if file_path.endswith(".pdf"):
                        text = self.extract_text_from_pdf(file_path)
                    elif file_path.endswith(".docx"):
                        text = self.extract_docx_text(file_path)
                    else:
                        continue

                    if not text.strip():
                        continue

                    resume_flag = self.is_resume(text)
                    att.is_resume = resume_flag
                    db.commit()

                    if not resume_flag:
                        if os.path.exists(file_path):
                            os.remove(file_path)
                        continue

                    logger.info("Extracting candidate info from attachment: %s", att.file_name)
                    info = self.extract_basic_info(text)
                    if info:
                        info["file_name"] = att.file_name
                        results.append({
                            "info": info,
                            "text": text,
                            "attachment_id": att.attachment_id,
                            "sender_email": parse_email_address(email_obj.sender),
                        })
                    kept_files.append((att.attachment_id, file_path))

                except Exception as e:
                    db.rollback()
                    logger.warning("Failed to process attachment %s: %s", att.file_name, e)
                    continue

            if results:
                self.save_resumes_to_db(results)

            for att_id, path in kept_files:
                if not os.path.exists(path):
                    continue

                compressed_name = None
                compressed_path = None
                try:
                    res_name = compress_file(path)
                    if (
                        res_name
                        and isinstance(res_name, str)
                        and not res_name.startswith("File does not exist")
                        and not res_name.startswith("Unsupported")
                        and "failed" not in res_name.lower()
                    ):
                        c_path = os.path.join(os.path.dirname(path), res_name)
                        if os.path.exists(c_path):
                            compressed_name = res_name
                            compressed_path = c_path
                except Exception as comp_err:
                    logger.warning("Compression failed for %s: %s; keeping original file", path, comp_err)

                # Target file to keep locally and upload
                target_path = compressed_path if compressed_path and os.path.exists(compressed_path) else path
                target_name = os.path.basename(target_path)

                try:
                    stored_name = attachment_storage_service.upload_attachment(target_path, target_name)
                except AttachmentStorageError as exc:
                    logger.warning("Storage upload failed for %s: %s; keeping local copy", target_path, exc)
                    stored_name = target_name

                attachment = db.query(Attachment).filter(Attachment.attachment_id == att_id).first()
                if attachment:
                    attachment.file_name = stored_name
                    db.commit()
                    db.refresh(attachment)

                # Remove only the duplicate uncompressed original if a distinct compressed copy exists
                if compressed_path and os.path.exists(compressed_path) and path != compressed_path:
                    if os.path.exists(path):
                        os.remove(path)
                # target_path is ALWAYS kept in the local directory

            return {
                "message_id": str(email_id),
                "processed_files": len(results),
            }


# =============================================================================
# 5. Candidate Document Service
# =============================================================================

class CandidateDocumentService(BaseService):
    """Encapsulates resume tailoring, document previews, batch downloads, docx and pdf generation."""

    @staticmethod
    def get_file_path(filename: str) -> str:
        """Resolve full path to an attachment file."""
        try:
            path = os.path.join(FILES_DIR, filename)
            if not os.path.exists(path):
                raise FileNotFoundError(f"{filename} not found")
            return path
        except FileNotFoundError:
            raise
        except Exception as e:
            logger.error("Error in get_file_path: %s", str(e))
            raise Exception("An error occurred while retrieving the file path")

    def get_file_for_download(self, filename: str) -> str:
        """Return path for downloading a file."""
        try:
            return self.get_file_path(filename)
        except FileNotFoundError as e:
            raise e
        except Exception as e:
            logger.error("Error in get_file_for_download: %s", str(e))
            raise HTTPException(status_code=404, detail=str(e))

    def get_file_base64(self, resume_id: Any) -> dict:
        """Fetch attachment file content encoded in base64 for in-browser preview."""
        if not resume_id or not str(resume_id).strip():
            raise HTTPException(status_code=400, detail="Resume ID is required")
        resume_id = str(resume_id).strip()
        with self.get_session() as db:
            try:
                resume_details = db.query(Resume).filter_by(resume_id=resume_id).first()
                if not resume_details:
                    raise HTTPException(status_code=404, detail="Resume not found")

                attachment_id = resume_details.attachment_id
                if not attachment_id:
                    raise HTTPException(status_code=404, detail="Attachment not found")

                attachment_details = db.query(Attachment).filter_by(attachment_id=attachment_id).first()
                if not attachment_details or not attachment_details.file_name:
                    raise HTTPException(status_code=404, detail="File name not found")

                filename = attachment_details.file_name
                raw_bytes = attachment_storage_service.read_attachment(filename, FILES_DIR)
                encoded = base64.b64encode(raw_bytes).decode()

                mime_type, _ = mimetypes.guess_type(filename)
                if mime_type != "application/pdf":
                    mime_type = "application/docx"

                return {
                    "status": status.HTTP_200_OK,
                    "message": "File previewed successfully",
                    "data": {
                        "name": filename,
                        "type": mime_type,
                        "content": encoded,
                    },
                }
            except HTTPException:
                raise
            except Exception as e:
                logger.error("Error in get_file_base64: %s", str(e), exc_info=True)
                raise HTTPException(status_code=404, detail=str(e))

    def get_download_links(self, candidate_ids: List[str], db: Optional[Session] = None) -> dict:
        """Create a zip archive of resumes for specified candidates and return base64 payload."""
        with self.get_session(db) as session:
            try:
                candidate_details = (
                    session.query(
                        Candidate.name,
                        Candidate.candidate_id,
                        Resume.resume_id,
                        Attachment.attachment_id,
                        Attachment.file_name,
                    )
                    .join(Resume, Candidate.candidate_id == Resume.candidate_id)
                    .join(Attachment, Resume.attachment_id == Attachment.attachment_id)
                    .filter(Resume.candidate_id.in_(candidate_ids))
                    .all()
                )

                if not candidate_details:
                    raise HTTPException(status_code=404, detail="No resumes found")

                folder_name = f"{uuid4()}_{datetime.datetime.now():%Y%m%d_%H%M%S}"
                download_path = os.path.join(BASE_DIR, DOWNLOAD_FILES)
                os.makedirs(download_path, exist_ok=True)

                temp_folder = os.path.join(download_path, folder_name)
                os.makedirs(temp_folder, exist_ok=True)

                for resume in candidate_details:
                    candidate_name = re.sub(r'[<>:"/\\|?*]', "_", resume.name or "Unknown_Candidate")
                    candidate_folder = os.path.join(temp_folder, candidate_name)
                    os.makedirs(candidate_folder, exist_ok=True)

                    destination_file = os.path.join(
                        candidate_folder, f"{resume.attachment_id}_{resume.file_name}"
                    )
                    try:
                        with open(destination_file, "wb") as destination:
                            destination.write(attachment_storage_service.read_attachment(resume.file_name, FILES_DIR))
                    except FileNotFoundError:
                        logger.warning("File not found for attachment_id %s", resume.attachment_id)
                        continue

                zip_path = shutil.make_archive(temp_folder, "zip", temp_folder)
                shutil.rmtree(temp_folder, ignore_errors=True)

                with open(zip_path, "rb") as zip_file:
                    zip_base64 = base64.b64encode(zip_file.read()).decode("utf-8")

                return {
                    "file_name": os.path.basename(zip_path),
                    "content": zip_base64,
                }

            except HTTPException:
                raise
            except Exception as e:
                logger.exception("Error in get_download_links: %s", str(e))
                raise HTTPException(
                    status_code=500, detail="An error occurred while processing the download request"
                )

    @staticmethod
    def encode_file_to_base64(file_path: str) -> str:
        """Read file from path and return base64-encoded string."""
        with open(file_path, "rb") as file_obj:
            return base64.b64encode(file_obj.read()).decode("utf-8")

    @staticmethod
    def parse_llm_json(content: str) -> dict:
        """Clean and parse JSON returned by LLM."""
        cleaned = (
            content.strip()
            .replace("```json", "")
            .replace("```", "")
            .replace("\ufeff", "")
            .replace("fastapi_api |", "")
            .replace("postgres_db |", "")
        )

        start = cleaned.find("{")
        end = cleaned.rfind("}")
        if start == -1 or end == -1:
            raise ValueError("No valid JSON object found")

        json_str = cleaned[start : end + 1]
        try:
            return json.loads(json_str)
        except json.JSONDecodeError:
            repaired = re.sub(r",(\s*[}\]])", r"\1", json_str)
            try:
                return json.loads(repaired)
            except json.JSONDecodeError as e2:
                raise ValueError(f"Invalid JSON: {e2.msg}")

    @staticmethod
    def filter_skills_by_jd(data: dict, jd_text: Any) -> dict:
        """Filter technical skills to keep those appearing in the JD."""
        jd_lower = str(jd_text).lower()
        resume = data.get("tailored_resume", data)
        skills = resume.get("technical_skills", {})

        for category, skill_list in skills.items():
            if isinstance(skill_list, list):
                skills[category] = [s for s in skill_list if str(s).lower().strip() in jd_lower]

        existing = set()
        for cat_list in skills.values():
            if isinstance(cat_list, list):
                for s in cat_list:
                    existing.add(str(s).lower())

        gap = data.get("skill_gap_analysis", {})
        gap["existing_skills"] = list(existing)
        if gap.get("missing_skills"):
            gap["missing_skills"] = [s for s in gap["missing_skills"] if str(s).lower() not in existing]

        return data

    def tailor_resume(self, resume_text: str, jd_info: Any) -> dict:
        """Call LLM to adapt resume against target JD requirements."""
        prompt = """
    You are a Resume Tailoring Engine. Return one valid JSON object. No markdown.
    Use RESUME_TEXT and JD_INFO to create a tailored resume.
    """
        system_message = (
            "Resume tailoring engine. Return one JSON object using the exact JSON schema. "
            "Target role, skills, and responsibility themes come from JD_INFO only."
        )

        full_content = f"RESUME_TEXT:\n{resume_text[:12000]}\n\nJD_INFO:\n{str(jd_info)[:8000]}"
        data = llm_service.call_llm(
            prompt=prompt + "\n\n" + full_content,
            system_prompt=system_message,
            is_json=True,
            # max_tokens=2000,
        )
        return self.filter_skills_by_jd(data, jd_info)

    def alter_resume(self, payload: dict, resume_id: Any) -> dict:
        """Tailor candidate resume against JD and produce updated PDF or docx."""
        with self.get_session() as db:
            try:
                if not payload or not resume_id:
                    raise ValueError("Payload and resume_id are required.")

                job_description = (
                    payload.get("job_description")
                    or payload.get("jobDescription")
                    or payload.get("jd")
                    or []
                )

                resume = db.query(Resume).filter(Resume.resume_id == resume_id).first()
                if resume is None:
                    return {"status": status.HTTP_404_NOT_FOUND, "message": "Resume not found."}

                attachment_id = resume.attachment_id
                att = db.query(Attachment).filter(Attachment.attachment_id == attachment_id).first()
                if not att:
                    return {"status": status.HTTP_404_NOT_FOUND, "message": "Attachment not found."}

                file_path = f"attachments/{att.file_name}"
                if file_path.endswith(".pdf"):
                    text = CandidateParserService.extract_text_from_pdf(file_path)
                elif file_path.endswith(".docx"):
                    text = CandidateParserService.extract_docx_text(file_path)
                else:
                    return {"status": status.HTTP_400_BAD_REQUEST, "message": "Unsupported file format"}

                clean_jd = (
                    ", ".join(job_description)
                    if isinstance(job_description, list)
                    else clean_text(job_description)
                )
                if not clean_jd:
                    raise ValueError("job_description or jobDescription is required.")

                combined_text = clean_jd.strip()
                data = self.tailor_resume(text, combined_text)

                output_ext = ".pdf" if file_path.lower().endswith(".pdf") else ".docx"
                output_path = f"attachments/tailored_{att.file_name.rsplit('.', 1)[0]}{output_ext}"
                if output_ext == ".pdf":
                    make_resume = self.generate_resume_pdf(data, output_path)
                else:
                    make_resume = self.generate_resume_docx(data, output_path)

                if not isinstance(make_resume, str):
                    return make_resume

                return {
                    "message": "resume tailored successfully",
                    "file_name": att.file_name.rsplit(".", 1)[0] + output_ext,
                    "file_type": "pdf" if output_ext == ".pdf" else "docx",
                    "mime_type": (
                        "application/pdf"
                        if output_ext == ".pdf"
                        else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    ),
                    "content": self.encode_file_to_base64(make_resume),
                }

            except Exception as e:
                logger.error("Error in alter_resume: %s", str(e), exc_info=True)
                return {"status": status.HTTP_500_INTERNAL_SERVER_ERROR, "message": str(e)}

    def generate_resume_docx(self, data: dict, output_path: str) -> Union[str, dict]:
        """Generate tailored Word document from structured JSON data."""
        try:
            doc = Document()
            section = doc.sections[0]
            section.top_margin = Inches(0.6)
            section.bottom_margin = Inches(0.6)
            section.left_margin = Inches(0.7)
            section.right_margin = Inches(0.7)

            resume = data.get("tailored_resume", data)

            def safe_get(*keys, default=""):
                for key in keys:
                    value = resume.get(key)
                    if value not in [None, "", [], {}]:
                        return value
                return default

            name = clean_text(safe_get("name"))
            role = clean_text(safe_get("role"))
            phone = clean_text(safe_get("phone_number", "phoneNumber"))
            email = clean_text(safe_get("email"))
            location = clean_text(safe_get("location"))

            title = doc.add_paragraph()
            title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run = title.add_run(name)
            run.bold = True
            run.font.size = Pt(20)

            if role:
                role_para = doc.add_paragraph()
                role_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                run = role_para.add_run(role)
                run.italic = True
                run.font.size = Pt(12)

            contact_items = [item for item in [phone, email, location] if item]
            if contact_items:
                contact = doc.add_paragraph()
                contact.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                contact.add_run(" | ".join(contact_items))

            summary = clean_text(safe_get("summary"))
            if summary:
                add_heading(doc, "Professional Summary")
                para = doc.add_paragraph(summary)
                para.style.font.size = Pt(11)

            technical_skills = safe_get("technical_skills", "technicalSkills", default={})
            if isinstance(technical_skills, dict) and any(technical_skills.values()):
                add_heading(doc, "Technical Skills")
                for category, skills in technical_skills.items():
                    if not skills:
                        continue
                    category_name = category.replace("_", " ").title()
                    skills_text = (
                        ", ".join(clean_text(skill) for skill in skills if skill)
                        if isinstance(skills, list)
                        else clean_text(skills)
                    )
                    doc.add_paragraph(f"{category_name}: {skills_text}", style="List Bullet")

            work_experience = safe_get("work_experience", "workExperience", default=[])
            if isinstance(work_experience, list) and work_experience:
                valid_experience = []
                for exp in work_experience:
                    if not isinstance(exp, dict):
                        continue
                    job_title = clean_text(exp.get("role") or exp.get("jobTitle"))
                    company = clean_text(exp.get("company"))
                    project_name = clean_text(exp.get("project_name") or exp.get("projectName"))
                    duration = clean_text(
                        exp.get("duration")
                        or " - ".join(
                            part for part in [clean_text(exp.get("startDate")), clean_text(exp.get("endDate"))] if part
                        )
                    )
                    responsibilities = exp.get("responsibilities")
                    description = clean_text(exp.get("description"))

                    if any([job_title, company, project_name, duration, responsibilities, description]):
                        valid_experience.append({
                            "job_title": job_title,
                            "company": company,
                            "project_name": project_name,
                            "duration": duration,
                            "responsibilities": responsibilities,
                            "description": description,
                        })

                if valid_experience:
                    add_heading(doc, "Work Experience")

                for exp in valid_experience:
                    title_para = doc.add_paragraph()
                    heading_parts = [part for part in [exp["job_title"], exp["company"]] if part]
                    run = title_para.add_run(" | ".join(heading_parts))
                    run.bold = True
                    run.font.size = Pt(12)

                    if exp["project_name"]:
                        project_para = doc.add_paragraph()
                        run = project_para.add_run(exp["project_name"])
                        run.italic = True
                        run.font.size = Pt(10)

                    if exp["duration"]:
                        date_para = doc.add_paragraph()
                        run = date_para.add_run(exp["duration"])
                        run.italic = True
                        run.font.size = Pt(10)

                    responsibilities = exp["responsibilities"]
                    if isinstance(responsibilities, list):
                        for bullet in responsibilities:
                            b_text = clean_text(bullet)
                            if b_text:
                                p = doc.add_paragraph(b_text, style="List Bullet")
                                p.style.font.size = Pt(11)
                    elif responsibilities:
                        p = doc.add_paragraph(clean_text(responsibilities))
                        p.style.font.size = Pt(11)
                    elif exp["description"]:
                        for bullet in [b.strip() for b in exp["description"].split(".") if b.strip()]:
                            p = doc.add_paragraph(bullet, style="List Bullet")
                            p.style.font.size = Pt(11)

            education = safe_get("education", default=[])
            if isinstance(education, list) and education:
                valid_education = [e for e in education if isinstance(e, dict) and any(e.values())]
                if valid_education:
                    add_heading(doc, "Education")
                    for edu in valid_education:
                        degree = clean_text(edu.get("degree"))
                        field = clean_text(edu.get("field"))
                        institution = clean_text(edu.get("institution"))
                        percentage = clean_text(edu.get("cgpa") or edu.get("percentageCGPA"))
                        duration = clean_text(edu.get("duration"))

                        para = doc.add_paragraph()
                        title_text = " ".join(p for p in [degree, field] if p)
                        run = para.add_run(title_text)
                        run.bold = True
                        run.font.size = Pt(12)

                        edu_lines = []
                        if institution:
                            edu_lines.append(institution)
                        if duration:
                            edu_lines.append(duration)
                        if percentage:
                            edu_lines.append(f"Percentage/CGPA: {percentage}")
                        if edu_lines:
                            para.add_run("\n" + "\n".join(edu_lines))

            certifications = safe_get("certifications", default=[])
            if isinstance(certifications, list) and certifications:
                valid_certs = [c for c in certifications if isinstance(c, dict) and any(c.values())]
                if valid_certs:
                    add_heading(doc, "Certifications")
                    for cert in valid_certs:
                        para = doc.add_paragraph(style="List Bullet")
                        title_bits = [
                            clean_text(cert.get(k)) for k in ("name", "organization", "year") if clean_text(cert.get(k))
                        ]
                        run = para.add_run(" - ".join(title_bits))
                        run.bold = True
                        if clean_text(cert.get("details")):
                            para.add_run(f"\n{clean_text(cert.get('details'))}")

            achievements = safe_get("achievements", default=[])
            if isinstance(achievements, list) and achievements:
                cleaned = [clean_text(i) for i in achievements if clean_text(i)]
                if cleaned:
                    add_heading(doc, "Achievements")
                    for item in cleaned:
                        p = doc.add_paragraph(item, style="List Bullet")
                        p.style.font.size = Pt(11)

            languages = safe_get("languages", default=[])
            if isinstance(languages, list):
                lang_text = ", ".join(clean_text(i) for i in languages if clean_text(i))
                if lang_text:
                    add_heading(doc, "Languages")
                    p = doc.add_paragraph(lang_text)
                    p.style.font.size = Pt(11)

            doc.save(output_path)
            return output_path

        except Exception as e:
            logger.error("Error in generate_resume_docx: %s", str(e), exc_info=True)
            return {"status": status.HTTP_500_INTERNAL_SERVER_ERROR, "message": str(e)}

    def generate_resume_pdf(self, data: dict, output_path: str) -> Union[str, dict]:
        """Generate tailored PDF document using ReportLab."""
        try:
            styles = getSampleStyleSheet()
            story = []
            resume = data.get("tailored_resume", data)

            def safe_get(*keys, default=""):
                for key in keys:
                    value = resume.get(key)
                    if value not in [None, "", [], {}]:
                        return value
                return default

            name = clean_text(safe_get("name"))
            role = clean_text(safe_get("role"))
            email = clean_text(safe_get("email"))
            phone = clean_text(safe_get("phone_number", "phoneNumber"))
            location = clean_text(safe_get("location"))
            summary = clean_text(safe_get("summary"))

            center_style = ParagraphStyle("CenterStyle", parent=styles["Normal"], alignment=TA_CENTER)

            if name:
                story.append(Paragraph(f"<b>{name}</b>", styles["Title"]))
            if role:
                story.append(Paragraph(role, center_style))

            contact_info = [item for item in [phone, email, location] if item]
            if contact_info:
                story.append(Paragraph(" | ".join(contact_info), center_style))

            story.append(Spacer(1, 10))

            if summary:
                story.append(Paragraph("<b>Professional Summary</b>", styles["Heading2"]))
                story.append(Paragraph(summary, styles["Normal"]))
                story.append(Spacer(1, 8))

            technical_skills = safe_get("technical_skills", "technicalSkills", default={})
            if isinstance(technical_skills, dict) and any(technical_skills.values()):
                story.append(Paragraph("<b>Technical Skills</b>", styles["Heading2"]))
                for category, skills in technical_skills.items():
                    if not skills:
                        continue
                    skills_text = (
                        ", ".join(clean_text(s) for s in skills if s)
                        if isinstance(skills, list)
                        else clean_text(skills)
                    )
                    cat_name = category.replace("_", " ").title()
                    story.append(Paragraph(f"<b>{cat_name}:</b> {skills_text}", styles["Normal"]))
                story.append(Spacer(1, 8))

            work_experience = safe_get("work_experience", "workExperience", default=[])
            if isinstance(work_experience, list) and work_experience:
                story.append(Paragraph("<b>Work Experience</b>", styles["Heading2"]))
                for exp in work_experience:
                    if not isinstance(exp, dict):
                        continue
                    title = clean_text(exp.get("role") or exp.get("jobTitle"))
                    company = clean_text(exp.get("company"))
                    duration = clean_text(exp.get("duration"))
                    heading = " | ".join(item for item in [title, company] if item)

                    if heading:
                        story.append(Paragraph(f"<b>{heading}</b>", styles["Normal"]))
                    if duration:
                        story.append(Paragraph(duration, styles["Italic"]))

                    responsibilities = exp.get("responsibilities", [])
                    if isinstance(responsibilities, list):
                        bullets = [
                            ListItem(Paragraph(clean_text(r), styles["Normal"]))
                            for r in responsibilities
                            if clean_text(r)
                        ]
                        if bullets:
                            story.append(ListFlowable(bullets, bulletType="bullet"))
                    elif responsibilities:
                        story.append(Paragraph(clean_text(responsibilities), styles["Normal"]))
                    story.append(Spacer(1, 6))

            education = safe_get("education", default=[])
            if isinstance(education, list) and education:
                story.append(Paragraph("<b>Education</b>", styles["Heading2"]))
                for edu in education:
                    if not isinstance(edu, dict):
                        continue
                    degree = clean_text(edu.get("degree"))
                    field = clean_text(edu.get("field"))
                    institution = clean_text(edu.get("institution"))
                    duration = clean_text(edu.get("duration"))
                    cgpa = clean_text(edu.get("cgpa") or edu.get("percentageCGPA"))

                    title = " ".join(item for item in [degree, field] if item)
                    if title:
                        story.append(Paragraph(f"<b>{title}</b>", styles["Normal"]))
                    details = [d for d in [institution, duration, f"CGPA/Percentage: {cgpa}" if cgpa else ""] if d]
                    if details:
                        story.append(Paragraph("<br/>".join(details), styles["Normal"]))
                    story.append(Spacer(1, 6))

            certifications = safe_get("certifications", default=[])
            if isinstance(certifications, list) and certifications:
                valid_certs = [c for c in certifications if isinstance(c, dict) and any(c.values())]
                if valid_certs:
                    story.append(Paragraph("<b>Certifications</b>", styles["Heading2"]))
                    for cert in valid_certs:
                        line = " - ".join(
                            clean_text(cert.get(k)) for k in ("name", "organization", "year") if clean_text(cert.get(k))
                        )
                        if line:
                            story.append(Paragraph(line, styles["Normal"]))
                    story.append(Spacer(1, 8))

            pdf = SimpleDocTemplate(output_path, pagesize=A4)
            pdf.build(story)
            return output_path

        except Exception as e:
            logger.exception("Error generating PDF: %s", str(e))
            return {"status": 500, "message": str(e)}

    def upload_attachment_to_storage(
        self,
        attachment_id: Union[str, UUID],
        expires_in: int = 3600,
    ) -> Dict[str, Any]:
        """Upload one existing local attachment to Supabase Storage."""
        with self.get_session() as db:
            attachment = db.query(Attachment).filter_by(attachment_id=attachment_id).first()
            if not attachment:
                raise HTTPException(status_code=404, detail="Attachment record not found")
            if not attachment.file_name:
                raise HTTPException(status_code=400, detail="Attachment has no filename")

            filename = os.path.basename(attachment.file_name)
            local_path = os.path.join(os.getenv("ATTACHMENT_DIR", "attachments"), filename)
            if not os.path.isfile(local_path):
                raise HTTPException(status_code=404, detail=f"Local attachment not found: {filename}")

            try:
                storage_key = attachment_storage_service.upload_attachment(local_path, filename)
                prefix = os.getenv("SUPABASE_S3_PREFIX", "resumes").strip("/")
                return {
                    "status": 200,
                    "message": "Attachment uploaded to Supabase successfully",
                    "data": {
                        "attachment_id": str(attachment.attachment_id),
                        "file_name": attachment.file_name,
                        "storage_key": f"{prefix}/{storage_key}" if prefix else storage_key,
                        "url": attachment_storage_service.create_attachment_download_url(storage_key, expires_in),
                        "url_expires_in": expires_in,
                    },
                }
            except AttachmentStorageError as exc:
                logger.exception("Unable to upload existing attachment %s", attachment_id)
                raise HTTPException(status_code=502, detail=str(exc)) from exc

    def preview_file(self, resume_id: Optional[str] = None) -> Dict[str, Any]:
        """Fetch attachment file content encoded in base64 for in-browser preview."""
        return self.get_file_base64(resume_id)

    def download_multiple_resumes(self, candidate_ids: List[Any]) -> Dict[str, Any]:
        """Bulk download candidate resumes as a zip archive."""
        return self.get_download_links(candidate_ids)

    def modify_resume(self, payload: dict, resume_id: str) -> Any:
        """Modify or tailor candidate resume against job description."""
        try:
            return self.alter_resume(payload, resume_id)
        except Exception as e:
            return JSONResponse(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                content={"message": "An error occurred while modifying the resume.", "details": str(e)},
            )


# =============================================================================
# 6. Unified Facade Class
# =============================================================================

class CandidateService(
    CandidateProfileService,
    CandidateSearchService,
    CandidateParserService,
    CandidateDocumentService,
):
    """Unified candidate service facade combining all sub-domain capabilities."""

    def __init__(self, db: Optional[Session] = None):
        super().__init__(db=db)


# Global singleton instance for convenient and backward-compatible usage
candidate_service = CandidateService()


# =============================================================================
# 7. Backward-Compatible Function Exports
# =============================================================================

# --- Candidate Profile ---
def candidate_datails(
    page: int = 1,
    sort_by: Optional[str] = None,
    sort_type: Optional[str] = None,
) -> List[Dict[str, Any]]:
    return candidate_service.get_candidate_details(page, sort_by, sort_type)


def get_candidate_details(
    page: int = 1,
    sort_by: Optional[str] = None,
    sort_type: Optional[str] = None,
) -> List[Dict[str, Any]]:
    return candidate_service.get_candidate_details(page, sort_by, sort_type)


def get_candidate_basic_info(
    filter_by: Optional[str] = None,
    filter_value: Optional[str] = None,
    page: int = 1,
    page_size: int = 10,
    sort_by: Optional[str] = None,
    sort_order: Optional[str] = None,
) -> Dict[str, Any]:
    return candidate_service.get_candidate_basic_info(
        filter_by, filter_value, page, page_size, sort_by, sort_order
    )


def _validate_data(data: Any) -> None:
    return candidate_service._validate_data(data)


def _build_subqueries(db: Session):
    return candidate_service._build_subqueries(db)


def _base_candidate_query(db: Session, edu_sq, skills_sq, work_sq):
    return candidate_service._base_candidate_query(db, edu_sq, skills_sq, work_sq)


def _extract_results(rows, total_count: int) -> List[Dict[str, Any]]:
    return candidate_service._extract_results(rows, total_count)


def _sort_by_name(db, sort_type, candidate_education, candidate_skills, candidate_work_exp, per_page, offset, total_count):
    return candidate_service._sort_by_name(db, sort_type, candidate_education, candidate_skills, candidate_work_exp, per_page, offset, total_count)


def _sort_by_experience(db, sort_type, candidate_education, candidate_skills, candidate_work_exp, per_page, offset, total_count):
    return candidate_service._sort_by_experience(db, sort_type, candidate_education, candidate_skills, candidate_work_exp, per_page, offset, total_count)


def _sort_by_year(db, sort_type, candidate_skills, candidate_work_exp, per_page, offset, total_count):
    return candidate_service._sort_by_year(db, sort_type, candidate_skills, candidate_work_exp, per_page, offset, total_count)


def _sort_by_percentage(db, sort_type, candidate_education, candidate_skills, candidate_work_exp, per_page, offset, total_count):
    return candidate_service._sort_by_percentage(db, sort_type, candidate_education, candidate_skills, candidate_work_exp, per_page, offset, total_count)


# --- Resume Search & Filtering ---
def search_resumes(filters: dict, export: bool = False) -> Any:
    return candidate_service.search_resumes(filters, export=export)


def export_candidate(search_id: str, page: int, page_size: int, export: bool):
    return candidate_service.export_candidate(search_id, page, page_size, export=export)


def _execute_search(
    resolved_filters: Optional[dict],
    page: int,
    page_size: int,
    sort_by: Optional[str],
    sort_order: str,
    export: bool = False,
):
    return candidate_service.execute_search(
        resolved_filters, page, page_size, sort_by, sort_order, export=export
    )


def _store_search_session(
    search_id: str,
    raw_filters: dict,
    resolved_filters: dict,
    user_query: str,
) -> None:
    return candidate_service.store_search_session(
        search_id, raw_filters, resolved_filters, user_query
    )


def _load_search_session(search_id: str) -> Optional[dict]:
    return candidate_service.load_search_session(search_id)


def _keyword_filter_attempts(user_query: str) -> List[dict]:
    return candidate_service._keyword_filter_attempts(user_query)


def _fallback_extract_and_resolve(user_query: str) -> Tuple[dict, dict]:
    return candidate_service._fallback_extract_and_resolve(user_query)


def _extract_and_resolve(user_query: str) -> Tuple[dict, dict]:
    return candidate_service.extract_and_resolve(user_query)


def nl_search_initial(
    user_query: str,
    page: Optional[int] = 1,
    page_size: Optional[int] = 20,
    sort_by: Optional[str] = None,
    sort_order: Optional[str] = "asc",
    export: bool = False,
) -> dict:
    return candidate_service.nl_search_initial(
        user_query, page, page_size, sort_by, sort_order, export=export
    )


def nl_search_paginate(
    search_id: str,
    page: Optional[int] = 1,
    page_size: Optional[int] = 20,
    sort_by: Optional[str] = None,
    sort_order: Optional[str] = "asc",
    export: bool = False,
) -> dict:
    return candidate_service.nl_search_paginate(
        search_id, page, page_size, sort_by, sort_order, export=export
    )


def extract_filters_from_query(query: str) -> dict:
    return candidate_service.extract_filters_from_query(query)


def resolve_dynamic_filters(dynamic_filters: dict) -> dict:
    return candidate_service.resolve_dynamic_filters(dynamic_filters)


def merge_filters(standard: dict, dynamic: dict) -> dict:
    return candidate_service.merge_filters(standard, dynamic)


def get_master_data() -> dict:
    return candidate_service.get_master_data()


def apply_filters(candidates: list, filters: dict) -> list:
    return candidate_service.apply_filters(candidates, filters)


# --- Resume Parser ---
def process_resumes(email_id: UUID) -> dict:
    return candidate_service.process_resumes(email_id)


def is_resume(text: str) -> bool:
    return candidate_service.is_resume(text)


def extract_basic_info(resume_text: str) -> dict:
    return candidate_service.extract_basic_info(resume_text)


def save_resumes_to_db(resumes: List[dict]) -> None:
    return candidate_service.save_resumes_to_db(resumes)


def extract_text_from_pdf(path: str) -> str:
    return candidate_service.extract_text_from_pdf(path)


def extract_docx_text(path: str) -> str:
    return candidate_service.extract_docx_text(path)


def extract_text_from_docx(path: str) -> str:
    return candidate_service.extract_text_from_docx(path)


def run_llm_json(prompt: str, text: str, system: str = "You only return JSON", max_retries: int = 3):
    return candidate_service.run_llm_json(prompt, text, system=system, max_retries=max_retries)


def clean_json_response(content: Union[str, dict]) -> dict:
    return candidate_service.clean_json_response(content)


# --- Document & Download ---
def get_file_path(filename: str) -> str:
    return candidate_service.get_file_path(filename)


def get_file_base64(resume_id: Any) -> dict:
    return candidate_service.get_file_base64(resume_id)


def get_file_for_download(filename: str) -> str:
    return candidate_service.get_file_for_download(filename)


def get_download_links(db: Optional[Session], candidate_ids: List[str]):
    return candidate_service.get_download_links(candidate_ids, db=db)


def encode_file_to_base64(file_path: str) -> str:
    return candidate_service.encode_file_to_base64(file_path)


def alter_resume(payload: dict, resume_id: Any) -> dict:
    return candidate_service.alter_resume(payload, resume_id)


def parse_llm_json(content: str) -> dict:
    return candidate_service.parse_llm_json(content)


def filter_skills_by_jd(data: dict, jd_text: Any) -> dict:
    return candidate_service.filter_skills_by_jd(data, jd_text)


def tailor_resume(resume_text: str, jd_info: Any) -> dict:
    return candidate_service.tailor_resume(resume_text, jd_info)


def generate_resume_docx(data: dict, output_path: str) -> Union[str, dict]:
    return candidate_service.generate_resume_docx(data, output_path)


def generate_resume_pdf(data: dict, output_path: str) -> Union[str, dict]:
    return candidate_service.generate_resume_pdf(data, output_path)


def get_candidates(
    filters: Optional[dict] = None,
    page: int = 1,
    page_size: int = 20,
) -> List[Dict[str, Any]]:
    return candidate_service.get_candidates(filters, page, page_size)


def generate_dynamic_filters(body: dict) -> Dict[str, Any]:
    return candidate_service.generate_dynamic_filters(body)


def filter_resumes(
    filters: dict,
    search_id: str,
    sort_by: Optional[str] = None,
    sort_order: Optional[str] = None,
    page: int = 1,
    page_size: int = 20,
) -> dict:
    return candidate_service.filter_resumes(filters, search_id, sort_by, sort_order, page, page_size)


def profile_search(body: dict) -> Dict[str, Any]:
    return candidate_service.profile_search(body)


def get_filter_options() -> Dict[str, Any]:
    return candidate_service.get_filter_options()


def upload_attachment_to_storage(
    attachment_id: Union[str, UUID],
    expires_in: int = 3600,
) -> Dict[str, Any]:
    return candidate_service.upload_attachment_to_storage(attachment_id, expires_in)


def preview_file(resume_id: Optional[str] = None) -> Dict[str, Any]:
    return candidate_service.preview_file(resume_id)


def modify_resume(payload: dict, resume_id: str) -> Any:
    return candidate_service.modify_resume(payload, resume_id)


__all__ = [
    "BaseService",
    "CandidateServiceError",
    "CandidateProfileService",
    "CandidateSearchService",
    "CandidateParserService",
    "CandidateDocumentService",
    "CandidateService",
    "candidate_service",
    "candidate_datails",
    "get_candidate_details",
    "get_candidate_basic_info",
    "get_candidates",
    "generate_dynamic_filters",
    "filter_resumes",
    "profile_search",
    "get_filter_options",
    "upload_attachment_to_storage",
    "preview_file",
    "modify_resume",
    "export_candidate",
    "search_resumes",
    "_execute_search",
    "_store_search_session",
    "_load_search_session",
    "nl_search_initial",
    "nl_search_paginate",
    "extract_filters_from_query",
    "resolve_dynamic_filters",
    "merge_filters",
    "get_master_data",
    "apply_filters",
    "process_resumes",
    "is_resume",
    "extract_basic_info",
    "save_resumes_to_db",
    "extract_text_from_pdf",
    "extract_docx_text",
    "extract_text_from_docx",
    "run_llm_json",
    "clean_json_response",
    "normalize_experience",
    "parse_email_address",
    "get_file_path",
    "get_file_base64",
    "get_file_for_download",
    "get_download_links",
    "encode_file_to_base64",
    "alter_resume",
    "parse_llm_json",
    "filter_skills_by_jd",
    "tailor_resume",
    "generate_resume_docx",
    "generate_resume_pdf",
    "clean_text",
    "add_heading",
    "BASE_DIR",
    "EXPORT_PATH",
    "FILES_DIR",
    "DOWNLOAD_FILES",
]
